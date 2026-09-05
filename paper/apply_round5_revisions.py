"""Propagate the fourth-round audit's Yang-Zhang correction into the manuscript.

    python paper/apply_round5_revisions.py [--base BASE.docx] [--out OUT.docx]

WHAT WENT WRONG. The fourth-round audit found that the adopted Yang-Zhang ratio, reported as
1.309, was produced by a specification the manuscript never described: the NUMERATOR was built
from the unadjusted previous-session close (``yang_zhang`` had no parameter through which the
adopted series could reach it) while the matched DENOMINATOR was built from the
corporate-action-adjusted close. Sections 4.5, 6.5, 7.2, 9 and the abstract all asserted that one
definition was applied throughout. It was not.

The estimator now takes the previous close explicitly (``nepsevol.estimators.range_.yang_zhang``),
both sides of every Yang-Zhang comparison are built from the same series, and the 2x2 in
``scripts/26_robustness.py`` rebuilds the numerator per definition rather than holding it fixed.
The corrected figures, on the same 135,899 matched stock-days:

    1.288   row-matched only, same-session benchmark, unadjusted previous close
    1.273   + horizon correction (unadjusted throughout)  -- the audit's own recomputation
    1.280   + corporate-action-adjusted previous close on BOTH sides  -- ADOPTED
    1.309   the superseded MIXED-definition figure, which no consistent specification produces

This script updates the four prose locations that quoted 1.309 (or its interval), refreshes
Tables 7, 12 and 15 from the regenerated outputs, and records the correction in Section 9 --
because a number that has already been circulated has to be corrected in public, not quietly
replaced. It does not re-insert any table, so it is safe to run after rounds 3 and 4.
"""

from __future__ import annotations

import argparse
import pathlib
import sys

import pandas as pd

try:
    import docx
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

ROOT = pathlib.Path(__file__).resolve().parents[1]
TAB = ROOT / "output" / "tables"
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from apply_referee_revisions import find, set_text, fill_table  # noqa: E402
from apply_round3_revisions import grow_table_columns  # noqa: E402

SUPERSEDED = "1.309"


def load_numbers():
    est = pd.read_csv(TAB / "table33_estimator_ratios_bootstrap.csv").set_index("estimator")
    yz = est.loc["Yang-Zhang"]
    t48 = pd.read_csv(TAB / "table48_yang_zhang_horizon.csv")

    def cell(prefix, horizon):
        q = t48[t48.previous_close_definition.str.startswith(prefix)
                & (t48.horizon_matched == horizon)]
        return float(q.sd_ratio.iloc[0])

    return {
        "yz": f"{float(yz.sd_ratio):.3f}",
        "yz_ci": f"[{yz.lo95_twoway:.3f}, {yz.hi95_twoway:.3f}]",
        "yz_rows": f"{int(yz.n_matched_rows):,}",
        "yz_rowmatched": f"{cell('unadjusted', False):.3f}",
        "yz_horizon": f"{cell('unadjusted', True):.3f}",
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default=str(ROOT / "paper" /
                    "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"))
    ap.add_argument("--out", default=str(ROOT / "paper" /
                    "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"))
    a = ap.parse_args()
    n = load_numbers()
    doc = docx.Document(a.base)

    # ── prose: the four locations that quoted the mixed-definition figure ───────────────────
    edits = [
        # Abstract - Findings
        ("Yang-Zhang requires matching sample, horizon and a corporate-action-adjusted previous "
         "close simultaneously; doing so moves its ratio to the total-risk proxy to 1.309.",
         "Yang-Zhang requires matching sample, horizon and a corporate-action-adjusted previous "
         "close simultaneously, on both sides of the comparison; doing so moves its ratio to "
         "the total-risk proxy to {yz}."),
        # 6.5
        ("moves it to the adopted figure of 1.309 [1.285, 1.335] on 135,899 matched stock-days "
         "(Table 12), which is what Table 7 and Section 7.2 report.",
         "moves it to the adopted figure of {yz} {yz_ci} on {yz_rows} matched stock-days "
         "(Table 12), which is what Table 7 and Section 7.2 report. The previous close enters "
         "Yang-Zhang's own overnight term as well as the benchmark's close-to-close return, so "
         "each figure quoted here is computed with the same definition on both sides; an "
         "earlier revision applied the adjustment to the benchmark only and reported the "
         "resulting mixed-definition ratio of 1.309, which no single specification produces."),
        # 7.2
        ("it stands at 1.309 [1.285, 1.335] of that proxy on 135,899 matched stock-days "
         "(Table 12)",
         "it stands at {yz} {yz_ci} of that proxy on {yz_rows} matched stock-days (Table 12)"),
        # 9 - the disclosed progression
        ("which moved it to the adopted 1.309.",
         "which moved it to {yz} once that definition was applied to Yang-Zhang's own overnight "
         "term as well as to the benchmark. A fourth-round audit found that the previously "
         "circulated 1.309 had applied the adjustment to the benchmark alone, because the "
         "estimator had no parameter through which the adopted previous close could reach it; "
         "the estimator now takes that series explicitly, a regression test fails if it is "
         "ignored, and the consistent alternatives are {yz_horizon} (unadjusted throughout) and "
         "{yz_rowmatched} (row-matched only, same-session benchmark)."),
    ]
    applied = 0
    for old, new in edits:
        paras = list(doc.paragraphs)
        hit = next((p for p in paras if old in p.text), None)
        if hit is None:
            print(f"  !! not found, skipped: {old[:70]!r}")
            continue
        set_text(hit, hit.text.replace(old, new.format(**n)))
        applied += 1
    print(f"  corrected {applied}/{len(edits)} prose locations")

    doc.save(a.out)

    # ── tables 7, 12, 15 refreshed from the regenerated outputs ─────────────────────────────
    doc = docx.Document(a.out)

    def refresh(caption, csv_name):
        paras = list(doc.paragraphs)
        cap_idx = find(paras, caption)
        children = list(doc.element.body)
        pos = children.index(paras[cap_idx]._p)
        n_before = sum(1 for el in children[:pos] if el.tag.endswith("}tbl"))
        # dtype=str: the CSVs already carry their intended display precision (e.g. "1.280");
        # re-parsing to float would silently render it as "1.28" in the manuscript table.
        t = pd.read_csv(TAB / csv_name, dtype=str).fillna("")
        tbl = doc.tables[n_before]
        grow_table_columns(tbl, len(t.columns))
        fill_table(tbl, list(t.columns), t.values.tolist())

    refresh("Table 7. All six named estimators", "paper_table7_all_estimators.csv")
    refresh("Table 12. Yang-Zhang against a horizon-matched benchmark",
            "paper_table12_yang_zhang_horizon.csv")
    refresh("Table 15. Dependence structure", "paper_table15_dependence_structure.csv")
    print("  refreshed Tables 7, 12, 15 from the corrected outputs")

    doc.save(a.out)
    print(f"\nwrote {a.out}")

    # ── verification ────────────────────────────────────────────────────────────────────────
    doc = docx.Document(a.out)
    paras = [p.text for p in doc.paragraphs]
    disclosure = "had applied the adjustment to the benchmark alone"
    body = "\n".join(p for p in paras if disclosure not in p and
                     "reported the resulting mixed-definition ratio" not in p)
    assert SUPERSEDED not in body, (
        f"the superseded {SUPERSEDED} survives outside the paragraphs that disclose it as "
        "superseded")
    assert any(disclosure in p for p in paras), "the correction disclosure is missing"
    assert n["yz"] in "\n".join(paras), f"the corrected figure {n['yz']} did not reach the text"
    # the table must agree with the text
    tbl_txt = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert n["yz"] in tbl_txt, "the corrected figure did not reach the tables"
    print(f"verification: {SUPERSEDED} appears only in its own correction disclosure; "
          f"{n['yz']} {n['yz_ci']} reached both text and tables")


if __name__ == "__main__":
    main()
