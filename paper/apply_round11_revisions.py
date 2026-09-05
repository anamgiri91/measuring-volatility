"""M7 final close: data-selected block length, exact interval wording, and origin counts.

    python paper/apply_round11_revisions.py [--base BASE.docx] [--out OUT.docx]

THREE CORRECTIONS.

1. THE MANUSCRIPT STILL CALLED THE ORIGINS "INDEPENDENT". They are not -- lag-1 autocorrelation
   0.646 -- and this script's own dependence analysis had already established that. The word is
   removed: the windows are non-overlapping, which is a statement about the outcome windows, not
   about the origins.

2. ONE PRIMARY BLOCK LENGTH, SELECTED FROM THE DATA. Reporting blocks of 1, 3 and 6 and then
   quoting one of them leaves the choice open to hindsight. The primary is now the automatic
   Politis-White (2004) / Patton-Politis-White (2009) selection, computed on each series and
   taken at the maximum: 10.88 origins. The grid is retained as sensitivity. The number of
   origins (190) and the implied effective-block counts are reported so a reader can judge how
   much independent information a block of that length leaves -- about 17 blocks, which is
   candidly not many, and is why the sensitivity grid is shown rather than hidden.

3. "DISTRIBUTION-FREE" IS REPLACED BY THE EXACT DESCRIPTION. A nonparametric bootstrap avoids a
   parametric normality assumption but is not literally distribution-free in finite samples. The
   Fisher-scale interval is now described as "a percentile interval computed from paired
   stationary-block bootstrap replicates after Fisher-z transformation", which is what it is.
"""

from __future__ import annotations

import argparse
import pathlib
import sys

import pandas as pd

try:
    import docx
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

ROOT = pathlib.Path(__file__).resolve().parents[1]
TAB = ROOT / "output" / "tables"
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from apply_referee_revisions import set_text  # noqa: E402


def load_numbers():
    bkall = pd.read_csv(TAB / "table68_forward_block_bootstrap.csv")
    bk = bkall[bkall.is_primary].set_index("estimator")
    db = pd.read_csv(TAB / "table69_forward_difference_block.csv")
    prim = db[db.is_primary].iloc[0]
    bls = pd.read_csv(TAB / "table70_block_length_selection.csv")
    cc, pk = bk.loc["Close-to-close"], bk.loc["Parkinson"]
    return {
        "b": f"{float(prim.mean_block_origins):.2f}",
        "n_org": int(prim.n_origins),
        "eff": f"{float(prim.effective_blocks):.0f}",
        "cc_ci": f"[{cc.corr_lo95:.3f}, {cc.corr_hi95:.3f}]",
        "pk_ci": f"[{pk.corr_lo95:.3f}, {pk.corr_hi95:.3f}]",
        "cc_slope_ci": f"[{cc.slope_lo95:.3f}, {cc.slope_hi95:.3f}]",
        "pk_slope_ci": f"[{pk.slope_lo95:.3f}, {pk.slope_hi95:.3f}]",
        "diff_ci": f"[{prim.lo95:.3f}, {prim.hi95:.3f}]",
        "z_ci": f"[{prim.fisher_z_lo95:.3f}, {prim.fisher_z_hi95:.3f}]",
        "grid": ", ".join(f"{b:g}" for b in sorted(db.mean_block_origins)),
        "sel_max": bls[bls.series.str.startswith("PRIMARY")].politis_white_block_length.iloc[0],
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default=str(ROOT / "paper" /
                    "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"))
    ap.add_argument("--out", default=str(ROOT / "paper" /
                    "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"))
    a = ap.parse_args()
    n = load_numbers()
    doc = docx.Document(a.base)
    edits = 0

    # 1 + 2. the "independent origins" claim, and the primary interval sentence
    for p in list(doc.paragraphs):
        if "independent forecast origins between" in p.text:
            set_text(p, p.text.replace(
                f"giving {n['n_org']} independent forecast origins between",
                f"giving {n['n_org']} forecast origins between"))
            edits += 1
            print("  removed the 'independent origins' claim")
            break

    old_ci = ("where the intervals come from a paired stationary block bootstrap rather than an "
              "analytic formula")
    new_ci = ("where the intervals come from a paired stationary block bootstrap rather than an "
              "analytic formula, at a mean block length of {b} origins selected automatically "
              "from the data by the Politis-White procedure rather than chosen by hand (blocks "
              "of {grid} origins are reported as sensitivity in Table 28, and the verdict is "
              "the same at every one). With {n_org} origins a block of that length leaves about "
              "{eff} effectively independent blocks, which is not many, and is why the "
              "sensitivity grid is shown rather than a single interval").format(**n)
    for p in list(doc.paragraphs):
        if old_ci in p.text:
            set_text(p, p.text.replace(old_ci, new_ci))
            edits += 1
            print("  documented the automatic block-length selection and origin counts")
            break

    # the quoted correlation intervals must be the primary-block ones
    for p in list(doc.paragraphs):
        if "close-to-close correlates at 0.695" in p.text:
            t = p.text
            import re
            t = re.sub(r"close-to-close correlates at 0\.695 \[[^\]]*\]",
                       f"close-to-close correlates at 0.695 {n['cc_ci']}", t)
            t = re.sub(r"Parkinson at 0\.630 \[[^\]]*\]",
                       f"Parkinson at 0.630 {n['pk_ci']}", t)
            set_text(p, t)
            edits += 1
            print("  requoted the correlation intervals at the primary block length")
            break

    # the slope intervals likewise
    for p in list(doc.paragraphs):
        if "contains one; we therefore do not assert attenuation" in p.text:
            import re
            t = re.sub(r"whose slope interval is \[[^\]]*\]",
                       f"whose slope interval is {n['pk_slope_ci']}", p.text)
            t = re.sub(r"whose interval \[[^\]]*\] contains one",
                       f"whose interval {n['cc_slope_ci']} contains one", t)
            set_text(p, t)
            edits += 1
            print("  requoted the slope intervals at the primary block length")
            break

    # 3. the difference sentence: exact Fisher wording, primary interval
    for p in list(doc.paragraphs):
        if "Pairing is necessary but not sufficient" in p.text:
            import re
            t = re.sub(r"at mean block lengths of [^;]*; it contains zero at every one of them "
                       r"\([^)]*\)",
                       f"at a mean block length of {n['b']} origins, selected from the data; it "
                       f"contains zero there ({n['diff_ci']}) and at every length in the "
                       f"sensitivity grid", p.text)
            set_text(p, t)
            edits += 1
            print("  requoted the difference interval at the primary block length")
            break

    # the Fisher-z description: rewrite the whole clause rather than splicing inside it, and
    # drop "distribution-free" everywhere -- a nonparametric bootstrap avoids a parametric
    # normality assumption but is not literally distribution-free in finite samples.
    for p in list(doc.paragraphs):
        if "Fisher" in p.text and "also containing zero" in p.text:
            # NB: the interval text contains decimal points, so the wildcard must allow "."
            t = re.sub(
                r"A (?:Fisher-transformed version of the same paired draw|percentile interval "
                r"computed from paired stationary-block bootstrap replicates after Fisher-z "
                r"transformation) gives .*?, also containing zero\.",
                "A percentile interval computed from paired stationary-block bootstrap "
                f"replicates after Fisher-z transformation gives {n['z_ci']}, also containing "
                "zero.", p.text)
            t = t.replace("both distribution-free paired intervals",
                          "both paired bootstrap intervals, which avoid that assumption,")
            set_text(p, t)
            edits += 1
            print("  corrected the Fisher-z description and removed 'distribution-free'")
            break

    for p in list(doc.paragraphs):
        if "distribution-free" in p.text:
            set_text(p, p.text.replace("distribution-free ", "").replace("distribution-free", ""))
            edits += 1
            print("  removed a further 'distribution-free' mention")

    doc.save(a.out)
    print(f"\n  {edits} edits; wrote {a.out}")

    body = "\n".join(p.text for p in docx.Document(a.out).paragraphs)
    assert "independent forecast origins" not in body, "the 'independent origins' claim survives"
    assert "selected automatically" in body
    assert "after Fisher-z transformation" in body
    assert "distribution-free" not in body, "'distribution-free' survives in the manuscript"
    assert f"{n['n_org']} forecast origins" in body
    print("verification: origins not called independent, block length data-selected, "
          "Fisher wording exact")


if __name__ == "__main__":
    main()
