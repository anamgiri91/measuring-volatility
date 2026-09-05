"""Apply mandatory items 3, 5 and 11 of the fourth-round review to the manuscript.

    python paper/apply_round6_revisions.py [--base BASE.docx] [--out OUT.docx]

  M3  EQUIVALENCE FRAMING. The +/-5% margin is described as "prespecified" throughout. It is
      not: it was declared after prior results existed, and an earlier version of its own
      justification cited effects already observed in this project. It is now described as
      declared post hoc, applied consistently. Separately, textbook TOST at one-sided alpha=.05
      corresponds to a 90% interval, not the 95% intervals reported here -- so the criterion in
      use is CONSERVATIVE rather than standard, and is now named as such. New Table 18 reports
      every verdict across margins of +/-2.5%, 5% and 10% and at both confidence levels.

  M5  ANNUALISATION. Section 7.1 quoted "about 226 sessions per year under the Sunday-Thursday
      regime and about 257 under the Monday-Friday regime". The second figure annualises 99
      sessions observed over 0.39 of a year, and both regimes schedule five sessions per week,
      so the weekday reform cannot by itself change the annual count -- the gap is the holiday
      composition of a short partial window. The regime-specific claim is withdrawn and replaced
      with the whole-sample observed rate plus an explicit sensitivity against 252.

  M11 WEIGHTING. The primary estimand is a ratio of pooled STOCK-DAY sums, which weights
      securities by how much they traded. That was never stated, and equal-security aggregation
      existed only for Parkinson and Rogers-Satchell by quintile. New Tables 19 and 20 report,
      for all six estimators, the stock-day-weighted ratio beside the equal-security mean and
      median, the 5th-95th percentile of the security-level ratio, and the ratio by length of
      observed history.
"""

from __future__ import annotations

import argparse
import pathlib
import sys

import pandas as pd

try:
    import docx
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

ROOT = pathlib.Path(__file__).resolve().parents[1]
TAB = ROOT / "output" / "tables"
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from apply_referee_revisions import find, para_after, set_text, new_table_after  # noqa: E402
from apply_round4_revisions import set_repeat_header_row  # noqa: E402


def load_numbers():
    n = {}
    ann = pd.read_csv(TAB / "table46_annualisation_regimes.csv").set_index("regime")
    whole = ann.loc["WHOLE SAMPLE (the reported A)"]
    n["A"] = f"{float(whole.sessions_per_year):.1f}"
    n["A_sessions"] = int(whole.sessions)
    n["A_years"] = f"{float(whole.span_years):.2f}"
    monfri = ann.loc["Mon-Fri (from 2026-04-06)"]
    n["monfri_sessions"] = int(monfri.sessions)
    n["monfri_years"] = f"{float(monfri.span_years):.2f}"
    n["monfri_rate"] = f"{float(monfri.sessions_per_year):.0f}"
    sunthu = ann.loc["Sun-Thu (to 2026-04-05)"]
    n["sunthu_rate"] = f"{float(sunthu.sessions_per_year):.0f}"
    a = float(whole.sessions_per_year)
    n["var_infl"] = f"{100 * (252 / a - 1):.1f}"
    n["vol_infl"] = f"{100 * ((252 / a) ** 0.5 - 1):.1f}"

    w = pd.read_csv(TAB / "table52_weighting_all_estimators.csv")
    n["w_gap"] = f"{float(w.weighting_gap.max()):.3f}"
    pk = w[w.estimator == "Parkinson"].iloc[0]
    n["pk_pooled"] = f"{float(pk.stockday_weighted):.3f}"
    n["pk_eq_mean"] = f"{float(pk.equal_security_mean):.3f}"
    n["pk_p05"] = f"{float(pk.security_p05):.3f}"
    n["pk_p95"] = f"{float(pk.security_p95):.3f}"

    s = pd.read_csv(TAB / "table54_equivalence_margin_sensitivity.csv")
    mcols = [c for c in s.columns if c.startswith("verdict_at_")]
    n["flip"] = f"{int(s[mcols].nunique(axis=1).gt(1).sum())} of {len(s)}"
    return n


# Whole-paragraph rewrites, keyed on an anchor substring. Used where the replacement restates
# material that appears later in the SAME paragraph -- a substring edit there would leave the
# superseded sentences standing after the new ones.
REPLACE_WHOLE = [
    # ── M5: §7.1 withdraws the regime-specific annualisation factor ──────────────────────────
    ("where v̂ is the chosen daily variance estimator and A is the market's own annual session",
     "where v̂ is the chosen daily variance estimator and A is the market's own annual session "
     "count rather than an imported convention. For NEPSE over the study window the detected "
     "trading calendar yields A ≈ {A} genuine sessions per year ({A_sessions} sessions over "
     "{A_years} years), against the 252 a developed-market default would impose; annualising "
     "with 252 instead would inflate variance by {var_infl}% and volatility by {vol_infl}%. "
     "This is a SAMPLE-PERIOD OBSERVED RATE, not a NEPSE convention, and it spans the April "
     "2026 trading-week reform. We do not quote a regime-specific factor. The post-reform "
     "Monday-Friday window contains only {monfri_sessions} sessions over {monfri_years} of a "
     "year, so annualising it ({monfri_rate}/year, against {sunthu_rate}/year pre-reform) "
     "extrapolates from a partial year; and because both regimes schedule five sessions per "
     "week, the reform cannot by itself change the annual session count — the difference "
     "reflects the holiday composition of that short window rather than the trading week. A "
     "usable regime-specific factor requires a complete observed year or an official "
     "weekday-plus-holiday projection. The report should show close-to-close volatility beside "
     "the range measure and disclose the security's median trade count and zero-range rate, so "
     "a reader can see the conditions under which the number was produced."),
]

REPLACE = [
    # ── M3 + M11: §6.2's margin sentence ─────────────────────────────────────────────────────
    ("Every ratio is also classified against a prespecified +/-5% equivalence margin (Table 7) "
     "rather than read from whether its confidence interval happens to contain one, because a "
     "wider, less informative interval covers one more easily and must not be read as stronger "
     "support.",
     "Every ratio is also classified against a stated +/-5% equivalence margin (Table 7) rather "
     "than read from whether its confidence interval happens to contain one, because a wider, "
     "less informative interval covers one more easily and must not be read as stronger "
     "support. Two qualifications belong with that margin. It was declared after the estimates "
     "existed rather than preregistered, so it is applied consistently but carries no "
     "protection against having been chosen to suit them; Table 18 therefore reports every "
     "verdict at ±2.5%, ±5% and ±10%, and {flip} reported verdicts change somewhere across that "
     "grid. And because textbook two one-sided tests at α = .05 correspond to a 90% interval "
     "rather than the 95% intervals reported here, the criterion in use is conservative — it "
     "declares equivalence less often than TOST would — which Table 18 also shows by reporting "
     "both levels."),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default=str(ROOT / "paper" /
                    "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"))
    ap.add_argument("--out", default=str(ROOT / "paper" /
                    "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"))
    a = ap.parse_args()
    n = load_numbers()
    doc = docx.Document(a.base)

    whole_applied = 0
    for anchor, text in REPLACE_WHOLE:
        paras = list(doc.paragraphs)
        hit = next((p for p in paras if anchor in p.text), None)
        if hit is None:
            print(f"  !! whole-paragraph anchor not found: {anchor[:60]!r}")
            continue
        set_text(hit, text.format(**n))
        whole_applied += 1
    print(f"  rewrote {whole_applied}/{len(REPLACE_WHOLE)} whole paragraphs")

    applied = 0
    for old, new in REPLACE:
        paras = list(doc.paragraphs)
        hit = next((p for p in paras if old in p.text), None)
        if hit is None:
            print(f"  !! not found, skipped: {old[:70]!r}")
            continue
        set_text(hit, hit.text.replace(old, new.format(**n)))
        applied += 1
    print(f"  replaced {applied}/{len(REPLACE)} paragraphs")

    # ── M5: drop the superseded regime-factor sentence wherever it survives ──────────────────
    stale_regime = ("The figure differs across the schedule change: about 226 sessions per year "
                    "under the Sunday-Thursday regime and about 257 under the Monday-Friday "
                    "regime, which is a further reason to date the regime boundary explicitly "
                    "rather than annualize the whole sample at one rate. Where a regime-specific "
                    "factor is used it should be reported as such. ")
    paras = list(doc.paragraphs)
    for p in paras:
        if stale_regime in p.text:
            set_text(p, p.text.replace(stale_regime, ""))
            print("  removed the superseded regime-specific annualisation sentence")
            break

    # ── M3 + M11: remaining "prespecified" occurrences in the body ───────────────────────────
    fixed = 0
    for p in list(doc.paragraphs):
        if "prespecified" in p.text:
            set_text(p, p.text.replace("prespecified +/-5%", "stated (post hoc) +/-5%")
                              .replace("prespecified ±5%", "stated (post hoc) ±5%")
                              .replace("prespecified margin", "stated (post hoc) margin")
                              .replace("prespecified equivalence margin",
                                       "stated (post hoc) equivalence margin"))
            fixed += 1
    print(f"  corrected 'prespecified' in {fixed} paragraphs")

    # ── M11: state the estimand's weighting in §6.1, pointing at Tables 19-20 ────────────────
    paras = list(doc.paragraphs)
    anchor_txt = "Once the universe is restricted to ordinary equity"
    hit = next((p for p in paras if anchor_txt in p.text), None)
    if hit is not None and "ratio of pooled stock-day sums" not in hit.text:
        para_after(hit,
            "A note on what the headline ratio weights. Every ratio in this paper is a ratio of "
            "pooled stock-day sums, so a security contributes in proportion to how many sessions "
            "it traded: long-lived, actively traded securities carry more of it than thin or "
            "briefly listed ones. That is the right estimand for the practical question — across "
            "the trading this market actually did, how does the estimator compare with the "
            "matched proxy in aggregate scale — but it is not the only one, and it is not "
            "neutral. Table 19 therefore reports, for all six estimators, the stock-day-weighted "
            "ratio beside the equal-security mean and median and the 5th-95th percentile of the "
            "security-level ratio. The two weightings agree closely: no estimator's "
            f"equal-security mean departs from its stock-day-weighted ratio by more than "
            f"{n['w_gap']}. What the aggregate does conceal is the spread, which is wide for "
            f"every estimator — Parkinson's own security-level ratios run from {n['pk_p05']} to "
            f"{n['pk_p95']} between the 5th and 95th percentiles against a pooled {n['pk_pooled']} "
            "— and Table 20 shows the ratio by length of observed history, so a reader can check "
            "that the pooled figure is not an artifact of long-lived securities dominating the "
            "sum.", "Normal")
        print("  added the estimand/weighting paragraph to §6.1")

    doc.save(a.out)

    # ── new Tables 18, 19, 20 after Table 17 ────────────────────────────────────────────────
    doc = docx.Document(a.out)
    paras = list(doc.paragraphs)
    cap17 = next(p for p in paras if p.text.startswith("Table 17."))
    children = list(doc.element.body)
    pos = children.index(cap17._p)
    n_before = sum(1 for el in children[:pos] if el.tag.endswith("}tbl"))
    anchor_tbl = doc.tables[n_before]

    def spacer_after(tbl):
        el = OxmlElement("w:p")
        tbl._tbl.addnext(el)
        return Paragraph(el, tbl._parent)

    anchor = spacer_after(anchor_tbl)
    for caption, csv_name in [
        ("Table 18. Equivalence verdicts across the margin grid and both confidence levels. The "
         "margin is declared post hoc, so verdicts are reported at ±2.5%, ±5% and ±10%; the 95% "
         "interval is a conservative criterion and the 90% interval is textbook TOST.",
         "paper_table18_equivalence_margin_sensitivity.csv"),
        ("Table 19. What the headline estimand weights: stock-day-weighted ratio against "
         "equal-security aggregation and the security-level spread, for all six estimators.",
         "paper_table19_weighting_all_estimators.csv"),
        ("Table 20. SD ratio to the matched proxy by length of observed history (terciles of "
         "observed stock-days per security).",
         "paper_table20_ratio_by_history_stratum.csv"),
    ]:
        t = pd.read_csv(TAB / csv_name, dtype=str).fillna("")
        cap = para_after(anchor, caption, "Normal")
        new_t = new_table_after(cap, doc, list(t.columns), t.values.tolist())
        set_repeat_header_row(new_t)
        anchor = spacer_after(new_t)
    print("  added Tables 18, 19, 20")

    doc.save(a.out)
    print(f"\nwrote {a.out}")

    # ── verification ────────────────────────────────────────────────────────────────────────
    doc = docx.Document(a.out)
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "prespecified" not in body, "'prespecified' survives in the manuscript body"
    assert "about 257 under the Monday-Friday regime" not in body, \
        "the superseded regime-specific annualisation figure survives"
    for needed in ("ratio of pooled stock-day sums", "declared after the estimates existed",
                   "Table 18.", "Table 19.", "Table 20.", "SAMPLE-PERIOD OBSERVED RATE"):
        assert needed in body, f"missing after round 6: {needed!r}"
    assert len(doc.tables) == 20, f"expected 20 tables, found {len(doc.tables)}"
    print("verification: post-hoc margin language, withdrawn regime factor, stated estimand, "
          f"and Tables 18-20 all present ({len(doc.tables)} tables)")


if __name__ == "__main__":
    main()
