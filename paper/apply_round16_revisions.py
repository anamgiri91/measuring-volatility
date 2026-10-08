"""Round 16: the claims made for Anam's estimator, rechecked against the evidence.

    python paper/apply_round16_revisions.py [--base BASE.docx] [--out OUT.docx]

At the author's request every claim the round-15 text makes for the estimator was checked against the
frozen tables, and the results were stress-tested after the fact by ``scripts/44_anam_recheck.py``
(tables 117-120; POST HOC, not pre-registered). This round corrects what the check found
(``AUDIT-REGISTER.md`` M-020 to M-025) and adds the recheck as Tables 37 and 38:

  * pre-registration stated too broadly ("every analysis"), the design said to rest on two regimes
    "alone", the holdout said to be unread, and "pre-registered" used for plans that only the
    package's own history timestamps (M-020);
  * "beats" used where the evidence is "no significant difference", the S&P 500 result credited to the
    estimator where it coincides with overnight-squared plus Parkinson, the indices' wins over
    close-to-close credited to the estimator although classical range estimators share them, a
    five-session result stated without its horizon, the best of five post hoc calibrations quoted
    alone, and Morocco's full-form result stated without its fragility (M-021);
  * the level credited to the estimator when the calibration produces it (M-022);
  * causal and novelty wording: "determines", "validated", a novelty statement that Hansen and Lunde
    (2005) already partly cover, "the data say that trust should be measured" when fixing b at zero did
    better, "overreacts" generalised from b < 1, and two different b statistics set side by side (M-023);
  * a protocol for Nepal that recommended the estimator where the Nepal evidence favours
    close-to-close (M-024);
  * the recheck itself and the data assumptions it tested (M-025).

RULES ENFORCED HERE
  * Every number is interpolated from a frozen table, and every corrected claim is asserted against
    the tables before it is written. No frozen verdict changes.
  * The recheck is post hoc and labelled so wherever it is quoted.
  * Edits replace text inside the single run that holds it, so no paragraph loses its formatting.
"""
from __future__ import annotations

import argparse
import pathlib
import re
import sys

import pandas as pd

try:
    import docx
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

ROOT = pathlib.Path(__file__).resolve().parents[1]
TAB = ROOT / "output" / "tables"
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT / "src"))
from apply_referee_revisions import find, new_table_after, para_after  # noqa: E402
from apply_round4_revisions import set_repeat_header_row  # noqa: E402
from apply_round14_revisions import BANNED, ORDINALS, para_with, replace_in, spacer_after, table_of  # noqa: E402
from apply_round15_revisions import NAME, SUBTITLE, TITLE, VARIANT, abstract_words, tt  # noqa: E402
from apply_round15_revisions import load as load_round15  # noqa: E402
from nepsevol.estimators import anam as AN  # noqa: E402

DOCX = ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"
RANGE = ["P", "GK", "RS", "o2+P", "o2+GK", "YZ (daily form)"]
STABLE = [("NEPSE", "A2"), ("NIFTY50", "test half"), ("SP500", "test half"), ("DSE 2023-2026", "test half"),
          ("DSE 2009-2021", "test half"), ("Vietnam 2007-2020", "test half"), ("Morocco 2012-2026", "test half")]
NEW_MARKETS = ["DSE 2023-2026", "DSE 2009-2021", "Vietnam 2007-2020", "Morocco 2012-2026"]
WORDS = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six"}
ORD16 = ORDINALS + ["Sixteenth"]


def pct(v) -> str:
    return f"{100 * float(v):.1f}"


def load() -> dict:
    n = dict(load_round15())
    lev = pd.read_csv(TAB / "table117_anam_recheck_level.csv")
    inf = pd.read_csv(TAB / "table118_anam_recheck_inference.csv")
    mult = pd.read_csv(TAB / "table119_anam_recheck_multiplicity.csv")
    dat = pd.read_csv(TAB / "table120_anam_recheck_data.csv")
    f101 = pd.read_csv(TAB / "table101_anam_holdout_forecast.csv")
    f113 = pd.read_csv(TAB / "table113_anam_morocco_forecast.csv")
    c106 = pd.read_csv(TAB / "table106_anam_posthoc_calibration.csv")

    def lv(market, span, est, col):
        return float(lev[(lev.market == market) & (lev.span == span) & (lev.estimator == est)][col].iloc[0])

    def tq(market, span, win, est, ref, col="t_frozen_NW_h"):
        g = inf[(inf.market == market) & (inf.test_span == span) & (inf.window == win)
                & (inf.estimator == est) & (inf.reference == ref)]
        assert len(g) == 1, (market, span, win, est, ref)
        return float(g[col].iloc[0])

    def dv(market, span, check):
        g = dat[(dat.market == market) & (dat.span == span) & (dat.check == check)]
        assert len(g) == 1, (market, span, check)
        return float(g.value.iloc[0])

    # ── R1 level: the calibration, not the kernel ────────────────────────────────────────────
    assert abs(lv("NEPSE", "C", "Anam (calibrated, frozen T2)", "calibrated_ratio") - float(n["lev_C"])) < 5e-4
    dev = lambda col, ests: max(abs(lv(m, s, e, col) - 1) for m, s in STABLE for e in ests)
    n["cls_cal_max"] = pct(dev("calibrated_ratio", RANGE))
    n["anam_lag_max"] = pct(dev("calibrated_before_window_ratio", ["Anam"]))
    n["cls_lag_max"] = pct(dev("calibrated_before_window_ratio", RANGE))
    assert float(n["lev_dev"]) <= float(n["cls_cal_max"]) <= 5.0, "the classical estimators are close once calibrated"
    c_cls = [lv("NEPSE", "C", e, "calibrated_ratio") for e in RANGE]
    n["C_cls_lo"], n["C_cls_hi"] = f"{min(c_cls):.3f}", f"{max(c_cls):.3f}"
    assert float(n["lev_C"]) < min(c_cls), "after the reform the kernel's level is closer than any calibrated classical"

    # ── the S&P 500 and the indices ─────────────────────────────────────────────────────────
    sp = f101[(f101.market == "SP500")].set_index(["window", "estimator"])
    for w in (5, 21):
        assert abs(sp.loc[(w, "Anam"), "QLIKE"] - sp.loc[(w, "o2+P"), "QLIKE"]) < 1e-9, "Anam = o2+P on the S&P 500"
    n["sp_lower"] = WORDS[sum(sp.loc[(5, e), "QLIKE"] < sp.loc[(5, "Anam"), "QLIKE"] for e in RANGE)]
    assert dv("SP500", "test half", "share of sessions with trailing b clipped at 1") == 1.0
    n["sp_b"] = f"{dv('SP500', 'test half', 'pooled b, whole span, unclipped'):.2f}"
    assert float(n["sp_b"]) > 1
    n["nifty_b"] = f"{dv('NIFTY50', 'test half', 'pooled b, whole span, unclipped'):.2f}"
    for mk, key in (("NIFTY50", "nifty_cls_beat"), ("SP500", "sp_cls_beat")):
        g = f101[(f101.market == mk) & (f101.window == 5) & f101.estimator.isin(RANGE)]
        n[key] = WORDS[int(((g.t_vs_CC < -1.96) & (g.dQLIKE_vs_CC < 0)).sum())]

    # ── R2/R3 inference: loss function, lags, multiplicity ──────────────────────────────────
    MA = ("Morocco 2012-2026", "test half")
    n["MA_nw2"] = tt(tq(*MA, 5, "Anam", "CC", "t_NW_2h"))
    for col in ("t_NW_2h", "t_NW_4h", "t_nonoverlapping", "t_MSE_NW_h"):
        assert tq(*MA, 5, "Anam", "CC", col) > -1.96, f"Morocco F1 under {col}"
    holm = mult.set_index(["plan", "rule", "market"])["Holm p across plans"]
    n["MA_holm"] = f"{holm[('M18', 'F1', 'Morocco 2012-2026')]:.3f}"
    assert float(n["MA_holm"]) > 0.025
    beats = mult[mult["frozen verdict: beats"]]
    others = beats[~((beats.plan == "M18") & (beats.rule == "F1"))]
    assert (others["Holm p across plans"] < 0.025).all(), "every other frozen win survives Holm across the plans"
    assert (others["beats under t_NW_2h"] & others["beats under t_NW_4h"]).all(), "and longer lags"
    n["nifty_nono"] = tt(tq("NIFTY50", "test half", 5, "Anam", "CC", "t_nonoverlapping"))
    assert float(n["nifty_nono"].replace("−", "-")) > -1.96
    assert set(beats[~beats["beats under t_nonoverlapping"]].rule + " " + beats[~beats["beats under t_nonoverlapping"]].market) \
        == {"F1 Morocco 2012-2026", "H4 NIFTY50"}, "only these two frozen wins fail with non-overlapping origins"
    # the classical range estimators never beat Anam under QLIKE, whatever the inference
    rr = inf[inf.estimator.isin(RANGE) & (inf.reference == "Anam")]
    for col in ("t_frozen_NW_h", "t_NW_2h", "t_NW_4h", "t_nonoverlapping"):
        assert (rr[col].dropna() > -1.96).all(), f"a classical range estimator beats Anam under {col}"
    OVERNIGHT = ["o2+P", "o2+GK", "YZ (daily form)"]
    wins = rr[rr["t_MSE_NW_h"] < -1.96]
    assert set(zip(wins.market, wins.window)) == {("Vietnam 2007-2020", 21)} and set(wins.estimator) == set(OVERNIGHT), \
        "under MSE only the estimators with a full overnight term beat Anam, and only in Vietnam at 21 sessions"
    n["VN21_mse_lo"], n["VN21_mse_hi"] = tt(wins["t_MSE_NW_h"].min()), tt(wins["t_MSE_NW_h"].max())
    VN = ("Vietnam 2007-2020", "test half")
    n["VN5_mse"] = tt(tq(*VN, 5, "Anam", "CC", "t_MSE_NW_h"))
    n["VN5_mse_of"] = tt(tq(*VN, 5, VARIANT, "CC", "t_MSE_NW_h"))
    n["A2C5_mse"] = tt(tq("NEPSE", "A2+C", 5, "Anam", "CC", "t_MSE_NW_h"))
    n["A2C5_mse_of"] = tt(tq("NEPSE", "A2+C", 5, VARIANT, "CC", "t_MSE_NW_h"))
    n["C5_mse"] = tt(tq("NEPSE", "C", 5, "Anam", "CC", "t_MSE_NW_h"))
    assert tq(*VN, 5, "Anam", "CC", "t_MSE_NW_h") > 1.96 and tq(*VN, 5, VARIANT, "CC", "t_MSE_NW_h") > 1.96
    assert tq("NEPSE", "A2+C", 5, "Anam", "CC", "t_MSE_NW_h") < -1.96 and tq("NEPSE", "A2+C", 5, VARIANT, "CC", "t_MSE_NW_h") < -1.96
    assert tq("SP500", "test half", 5, "Anam", "CC", "t_MSE_NW_h") < -1.96
    assert abs(tq("NEPSE", "C", 5, "Anam", "CC", "t_MSE_NW_h")) < 1.96

    # ── Morocco by band regime (reported in table113, not decided) ──────────────────────────
    reg = f113[(f113.window == 5) & (f113.estimator == "Anam") & (f113.test_span != "test half")].set_index("test_span")["t_vs_CC"]
    k4 = next(s for s in reg.index if s.startswith("4%"))
    n["MA_4pct"] = tt(reg[k4])
    rest = reg.drop(k4)
    n["MA_rest_lo"], n["MA_rest_hi"] = tt(rest.min()), tt(rest.max())
    assert reg[k4] < -1.96 and (rest.abs() < 1.96).all(), "Morocco's full-form win sits in the 4% regime alone"

    # ── the post hoc calibration variants of regime C (table106) ────────────────────────────
    anam_rows = c106[c106.estimator.str.startswith("Anam")]
    n["n_windows"] = WORDS[anam_rows.estimator.str.extract(r"calibration (\d+) dates")[0].dropna().nunique()]
    assert (anam_rows.dQLIKE_vs_CC > 0).all(), "close-to-close kept the lower loss under every variant"
    best = anam_rows[anam_rows.window == 5].sort_values("QLIKE").iloc[0]
    assert best.estimator == "Anam, calibration 10 dates", "the ten-date window is the one quoted"

    # ── data checks (table120) ──────────────────────────────────────────────────────────────
    n["ma_zero"] = f"{100 * dv('Morocco 2012-2026', 'test half', 'zero-range share (high = low)'):.0f}"
    bs = [dv(m, "test half", "pooled b, whole span, unclipped") for m in NEW_MARKETS]
    n["b_lo"], n["b_hi"] = f"{min(bs):.3f}", f"{max(bs):.3f}"
    for span, key in (("A2", "b_nepse_hi"), ("C", "b_nepse_lo")):
        assert f"{dv('NEPSE', span, 'pooled b, whole span, unclipped'):.3f}" == str(n[key]), "same statistic as M15"
    R, U = "day <= 12 (repaired in 2009-2021)", "day > 12 (stamped correctly)"
    D9 = "DSE 2009-2021"
    n["rep_big"] = pct(dv(D9, R, "date repair: share |r| > 0.05"))
    n["unr_big"] = pct(dv(D9, U, "date repair: share |r| > 0.05"))
    n["rep_b"] = f"{dv(D9, R, 'date repair: pooled b, unclipped'):.2f}"
    n["unr_b"] = f"{dv(D9, U, 'date repair: pooled b, unclipped'):.2f}"
    assert dv(D9, R, "date repair: dates on a Friday or Saturday") == 0
    n["dse_close_diff"] = f"{100 * (1 - dv('DSE 2023-2026', 'mirror, 2024-09-30 to 2026-10-08', 'close equals last-trade price (share)')):.0f}"
    q = dat[dat.check == "band schedule: 99.9th percentile of |r| (no band screen)"].set_index("span").value
    order = ["10% to 2020-03-16", "4% 2020-03-17 to 2021-10-11", "6% 2021-10-12 to 2023-10-08", "10% from 2023-10-09"]
    for s, lim in zip(order, (0.10, 0.04, 0.06, 0.10)):
        assert lim - 0.005 <= q[s] <= lim + 0.01, f"the largest moves do not sit at the {s} limit"
    n["ma_q"] = ", ".join(f"{100 * q[s]:.1f}%" for s in order[:-1]) + f" and {100 * q[order[-1]]:.1f}%"
    vb = dat[dat.check == "exchange mix: pooled b, unclipped"].value
    n["vn_b_lo"], n["vn_b_hi"] = f"{vb.min():.2f}", f"{vb.max():.2f}"
    assert vb.max() - vb.min() < 0.1
    n["pool"] = AN.POOL_SESSIONS
    return n


# ─────────────────────────────────────────────────────────────────── docx utilities

def edit(p, old: str, new: str) -> None:
    """Replace ``old`` inside the one run that holds it, so the paragraph's other runs keep their format."""
    hits = [r for r in p.runs if old in r.text]
    assert len(hits) == 1 and hits[0].text.count(old) == 1, f"expected {old[:70]!r} once, in one run, of {p.text[:60]!r}"
    hits[0].text = hits[0].text.replace(old, new)


# ─────────────────────────────────────────────────────────────────── the revision

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default=str(DOCX))
    ap.add_argument("--out", default=str(DOCX))
    a = ap.parse_args()
    n = load()
    doc = docx.Document(a.base)
    if any(p.text.startswith("Table 37.") for p in doc.paragraphs):
        sys.exit("round 16 has already been applied to this document")
    if not any(p.text.startswith("6.8 ") for p in doc.paragraphs):
        sys.exit("round 15 must be applied first")

    p = para_with(doc, "Data and reproducibility note.")
    replace_in(p, "through apply_round15_revisions.py)", "through apply_round16_revisions.py)")

    # ── abstract (M-020, M-021, M-022, M-023) ───────────────────────────────────────────────
    edit(para_with(doc, "Design/methodology/approach: "), "Every analysis follows a plan frozen before testing.",
         "Later analyses follow plans frozen before testing.")
    edit(para_with(doc, "Findings: "),
         f"No classical range estimator beats it in any test sample, and its level is within {n['lev_dev']}% of "
         "close-to-close outside NEPSE's post-reform regime, but it does not beat close-to-close at short horizons "
         "in Nepal, Bangladesh or Vietnam. Its open-free form passed a pre-registered test in Morocco.",
         "No classical range estimator forecasts significantly better than it in any test sample, but it does not "
         "beat close-to-close at short horizons in Nepal, Vietnam or recent Dhaka data, and both comparisons "
         "depend partly on the loss function. Its open-free form passed a prespecified test in Morocco.")
    edit(para_with(doc, "Originality/value: "),
         "A market-design rule determines what daily bars measure, and an estimator built on that evidence is "
         "validated out of sample in four frontier markets.",
         "A market-design rule change altered what daily bars measure; an estimator built on that evidence is "
         "tested out of sample in four frontier markets.")
    words = abstract_words(doc)
    assert words <= 255, f"abstract is {words} words"
    print(f"  corrected the abstract ({words} words with its labels)")

    # ── introduction ─────────────────────────────────────────────────────────────────────────
    p = para_with(doc, "The fifth answer builds on the fourth.")
    edit(p, "It was designed on NEPSE's first two rule regimes alone and then evaluated",
         "It was designed on NEPSE's first two rule regimes, with Section 6.7's findings for the whole NEPSE "
         "sample already known, and then evaluated")
    edit(p, f"No classical range estimator beats it in any of the seven test samples, and its calibrated level "
            f"stays within {n['lev_dev']}% of close-to-close variance everywhere except NEPSE's post-reform "
            f"regime, while the classical estimators overstate it by up to {n['over_max']}% and understate it by "
            f"up to {n['under_max']}% on the same samples.",
         "Under the loss function the plans fixed, no classical range estimator has significantly lower loss than "
         f"it in any of the seven test samples. Its calibrated level stays within {n['lev_dev']}% of close-to-close "
         "variance everywhere except NEPSE's post-reform regime; uncalibrated, the classical estimators "
         f"overstate that level by up to {n['over_max']}% and understate it by up to {n['under_max']}% on the "
         f"same samples, but given the same calibration they come within {n['cls_cal_max']}%, so the level is "
         "the calibration's doing rather than the estimator's.")
    edit(p, "It does not beat close-to-close itself at short horizons in Nepal, Bangladesh or Vietnam,",
         "It does not beat close-to-close itself at short horizons in Nepal, Vietnam or the recent Dhaka panel,")
    edit(p, "and then passed a test fixed in advance in Morocco.",
         "and then passed a test fixed in advance in Morocco. A post hoc recheck finds the comparisons with "
         "the classical range estimators robust to the choice of standard errors, and those with "
         "close-to-close dependent on the loss function (Table 37).")
    print("  corrected the introduction")

    # ── Section 6.8 ──────────────────────────────────────────────────────────────────────────
    p = para_with(doc, "Sections 6.6 and 6.7 say what an estimator for this market needs.")
    edit(p, "What is added is that the open's weight is measured from the data rather than assumed to be one.",
         "Hansen and Lunde also weight the overnight return by an amount estimated from data, chosen for "
         "efficiency from intraday returns. What is added here is narrower: the weight is the open's "
         "unbiasedness coefficient for the close, measured from daily bars alone, and the same coefficient "
         "moves the range's anchor to the effective open and sets how much close-to-close variance is blended "
         "in.")
    p = para_with(doc, "The estimator was designed on NEPSE's first two rule regimes (A1 and B)")
    edit(p, "and its constants were fixed in a plan frozen before any later observation was read.",
         "and its constants were fixed in a plan frozen before the estimator was computed on any later "
         "observation. The later regimes were not unseen: Section 6.7's findings, which informed the design, "
         "use the whole NEPSE sample, and the NIFTY 50's unbiasedness coefficient was known; the plan lists "
         "what was known when it was frozen.")
    edit(p, "so ranking by QLIKE loss matches", "so ranking by expected QLIKE loss matches")
    edit(p, "not for any property of their prices.",
         "not for any property of their prices. Bangladesh, Vietnam and Morocco are frontier markets in MSCI's "
         "classification over the test periods, and FTSE Russell announced in October 2025 that it would "
         "reclassify Vietnam as secondary emerging from September 2026, after the Vietnamese sample ends; "
         "Nepal is not covered by the major index providers' frontier-market indices and is called a frontier "
         "market here in the descriptive sense.")
    p = para_with(doc, "Against the classical estimators the result is uniform.")
    edit(p, "Against the classical estimators the result is uniform. No classical range estimator beats "
            f"{NAME} in any of the seven test samples at either horizon (Table 34), and in every "
            "frontier-market sample it beats all six at five sessions.",
         "Against the classical estimators the result is uniform under the plans' loss function. No classical "
         f"range estimator has significantly lower loss than {NAME} in any of the seven test samples at "
         "either horizon (Table 34), and in every frontier-market sample it beats all six at five sessions; "
         "a post hoc recheck finds the same with longer Newey-West lags and with non-overlapping forecast "
         "origins. Two qualifications apply. On the S&P 500 the index open lags the overnight move, with a "
         f"pooled unbiasedness coefficient of {n['sp_b']}, so the estimator's b sits at its cap of one in "
         "every test session and the estimator coincides with overnight² + Parkinson; "
         f"{n['sp_lower']} classical estimators had nominally lower loss there. And under an MSE loss, also "
         "post hoc, the three classical estimators that add the whole overnight move (overnight² + Parkinson, "
         "overnight² + Garman-Klass and the Yang-Zhang daily form) have significantly lower loss than the "
         f"estimator in Vietnam at 21 sessions (t from {n['VN21_mse_lo']} to {n['VN21_mse_hi']}).")
    edit(p, "a thin market in which many bars have no range at all.",
         f"a thin market in which {n['ma_zero']}% of the test-span bars have no range at all.")
    edit(p, "which is the case for setting the level on each market's own close-to-close scale.",
         "which is the case for setting the level on each market's own close-to-close scale. It is a case for "
         "the calibration, not for this kernel: given the same calibration, a post hoc check finds every "
         f"classical range estimator within {n['cls_cal_max']}% of close-to-close variance outside NEPSE's "
         "post-reform regime (Table 38). The level test is also partly in sample, because the calibration "
         "window contains the 21 sessions being measured; with a calibration that ends before them, the "
         f"estimator is within {n['anam_lag_max']}% and the classical estimators within {n['cls_lag_max']}%. "
         "The plans' level hypotheses (H2 and F3), which set a calibrated estimator against raw ones, are "
         "therefore weak tests. Only after NEPSE's reform, where every calibration lags, does the kernel "
         f"matter for the level: the estimator's {n['lev_C']} is closer to one than any calibrated classical "
         f"estimator's ({n['C_cls_lo']} to {n['C_cls_hi']}).")
    p = para_with(doc, "Against close-to-close the record is mixed, and frozen predictions failed.")
    edit(p, "where the plans had predicted that it would.",
         "where the plans had predicted that it would. On the indices it shares those wins with the classical "
         f"range estimators: {n['nifty_cls_beat']} of them also beat close-to-close on the NIFTY 50, and "
         f"{n['sp_cls_beat']} on the S&P 500. The Moroccan difference is the least secure. It comes "
         "from the 4% band regime of March 2020 to October 2021, which spans the pandemic shock (t = "
         f"{n['MA_4pct']} there, against {n['MA_rest_lo']} to {n['MA_rest_hi']} in the other three regimes), "
         "and in a post hoc recheck it does not survive longer Newey-West lags (t = "
         f"{n['MA_nw2']} with lags of twice the horizon), non-overlapping origins, an MSE loss or a Holm "
         f"correction across the three plans (p = {n['MA_holm']}); the NIFTY 50 difference does not survive "
         f"non-overlapping origins either (t = {n['nifty_nono']}) (Table 37).")
    edit(p, "in a post hoc check, a ten-date calibration lowers the estimator's five-session loss in that regime "
            f"from {n['y1_60']} to {n['y1_10']}, against {n['y1_cc']} for close-to-close, and the difference is "
            f"no longer significant (t = {n['y1_t10']}).",
         f"in a post hoc check of {n['n_windows']} calibration windows, the frozen {n['pool']} dates among them, "
         "and a reset at the rule change, the "
         f"ten-date window came closest, lowering the estimator's five-session loss in that regime from "
         f"{n['y1_60']} to {n['y1_10']}, against {n['y1_cc']} for close-to-close, so that the difference is no "
         f"longer significant (t = {n['y1_t10']}); close-to-close kept the lower loss under every variant tried.")
    p = para_with(doc, "The open-free form was reported, but not decided on, in the first two plans.")
    edit(p, "On the two indices, where the open is close to efficient, the full form had the lower loss,",
         "On the two indices, whose opens do not overreact (the pooled unbiasedness coefficient is "
         f"{n['nifty_b']} on the NIFTY 50 and above one on the S&P 500), the full form had the lower loss,")
    edit(p, "The evidence therefore supports the open-free form for markets whose open overreacts, with b well "
            "below one, as it is in every frontier market examined here (pooled medians between "
            f"{n['b_new_lo']} and {n['b_new_hi']} on the new test spans, against {n['b_nepse_lo']} to "
            f"{n['b_nepse_hi']} across NEPSE's regimes), and the full form where the open is reliable.",
         "The evidence therefore favours the open-free form over the full form where the open is unreliable, "
         f"with b well below one, as it is in every frontier market examined here ({n['b_lo']} to {n['b_hi']} "
         f"on the new test spans and {n['b_nepse_lo']} to {n['b_nepse_hi']} across NEPSE's regimes, each "
         "pooled over the whole span), and the full form where the open is close to unbiased. A coefficient "
         "below one shows how much of the overnight move the session reverses, not why: that the reversal is "
         "overreaction is Section 6.7's evidence for NEPSE alone, and in the other markets a noisy first "
         "trade, such as a bid-ask bounce, would lower b in the same way.")
    edit(p, f"(for the open-free form, t = {n['C5_of_t']} at five sessions in NEPSE's post-reform regime).",
         f"(for the open-free form, t = {n['C5_of_t']} at five sessions in NEPSE's post-reform regime). And "
         "the comparison with close-to-close depends on the loss function, which the plans fixed in advance as "
         "QLIKE: under an MSE loss, "
         f"a post hoc check, close-to-close beats both forms in Vietnam (t = {n['VN5_mse']} and "
         f"{n['VN5_mse_of']} at five sessions), both forms beat close-to-close in NEPSE's holdout (t = "
         f"{n['A2C5_mse']} and {n['A2C5_mse_of']}) and on the S&P 500, and close-to-close's advantage after "
         f"the reform is no longer significant (t = {n['C5_mse']} for the full form).")
    print("  corrected Section 6.8")

    # ── table captions ──────────────────────────────────────────────────────────────────────
    edit(para_with(doc, "Table 33. "), "the definition frozen in plan M16 before any holdout observation was read "
                                       "(Section 6.8).",
         "the definition frozen in plan M16 before the estimator was computed on any holdout sample (Section "
         "6.8; the plan lists what was already known about the holdout).")
    edit(para_with(doc, "Table 35. "), "and the classical estimators raw (Section 6.8).",
         "and the classical estimators raw (Section 6.8); Table 38 gives them the same calibration.")
    edit(para_with(doc, "Table 36. "), "Failed predictions are listed with the ones that held.",
         "Failed predictions are listed with the ones that held. 'Best in market' and 'best in panel' are the "
         "plans' names for 'no rival has significantly lower loss at either horizon'; they do not mean the "
         "lowest loss.")

    # ── Tables 37-38, after Table 36 ─────────────────────────────────────────────────────────
    doc.save(a.out)
    doc = docx.Document(a.out)
    anchor = spacer_after(table_of(doc, "Table 36."))
    for caption, csv_name in [
        ("Table 37. Post hoc robustness of the plans' comparisons (script 44; not pre-registered, run after "
         "every verdict was known): the frozen Newey-West t of each loss difference (negative favours the "
         "first estimator) against longer lags, non-overlapping forecast origins, an MSE loss on the same "
         "forecasts and each half of the test span, with the Holm-adjusted one-sided p of each plan rule "
         "across the three plans (a frozen 'beats' is p < 0.025).", "paper_table37_anam_robustness.csv"),
        (f"Table 38. Level under the same calibration (post hoc, script 44): each estimator's 21-session "
         f"mean, calibrated by the pooled scheme of {NAME}, divided by the 21-session mean of r², with the "
         "calibration ending at the window's last session as in Table 35 and, in the last two columns, "
         "ending before the window starts.", "paper_table38_anam_level_same_calibration.csv"),
    ]:
        t = pd.read_csv(TAB / csv_name, dtype=str).fillna("")
        cap = para_after(anchor, caption, "Normal")
        new_t = new_table_after(cap, doc, list(t.columns), t.values.tolist())
        set_repeat_header_row(new_t)
        anchor = spacer_after(new_t)
    print("  added Tables 37-38")

    # ── Section 7 (M-024) ────────────────────────────────────────────────────────────────────
    p = para_with(doc, f"Where one daily estimator must serve, use {NAME} (Section 6.8)")
    edit(p, f"Where one daily estimator must serve, use {NAME} (Section 6.8): its open-free form where the "
            "open's unbiasedness coefficient is well below one, as in every frontier market examined here, and "
            "its full form where the open is close to unbiased. Its calibrated level is on the close-to-close "
            "scale by construction. Report close-to-close beside it, and after a change in market rules rely "
            "on close-to-close until the estimator's calibration window has passed the change.",
         f"In NEPSE neither form of {NAME} (Section 6.8) beat close-to-close out of sample, and close-to-close "
         "beat both after the reform, so close-to-close remains the primary measure here. Where a range-based "
         "measure is wanted beside it, use the estimator's open-free form where the open's unbiasedness "
         "coefficient is well below one, as in every frontier market examined here, and its full form where "
         "the open is close to unbiased; its level is on the close-to-close scale by construction of its "
         "calibration. After a change in market rules, rely on close-to-close until the calibration window has "
         "passed the change.")
    p = para_with(doc, "where v̂ is the chosen daily variance estimator and A is the market's own annual")
    edit(p, "because κ already sets it on the close-to-close scale.",
         f"because κ already sets it on the close-to-close scale of its trailing {n['pool']} dates; after a "
         "change in market rules that scale lags (Section 6.8).")
    p = para_with(doc, "Proximity to a proxy is one criterion; it is not the only one,")
    edit(p, f"and under it no classical range estimator beats {NAME} in any market tested; even so it yields no "
            "universal ordering, because close-to-close is not beaten everywhere.",
         f"and under it no classical range estimator has significantly lower loss than {NAME} in any market "
         "tested; even so it yields no universal ordering, because close-to-close is not beaten everywhere and "
         "the comparison with it changes under a different loss function.")
    print("  corrected the protocol (Section 7), 7.1 and 7.2")

    # ── Discussion ───────────────────────────────────────────────────────────────────────────
    p = para_with(doc, "The estimator of Section 6.8 is the constructive side of these findings")
    edit(p, "Every classical formula fixes how far the open is trusted; the data say that trust should be "
            f"measured, because it varies across NEPSE's own rule regimes ({n['b_nepse_lo']} to "
            f"{n['b_nepse_hi']}) and across the frontier markets examined (pooled medians between "
            f"{n['b_new_lo']} and {n['b_new_hi']}).",
         "Each classical range formula compared here fixes how far the open is trusted, and that trust should "
         f"be checked rather than assumed, because it varies across NEPSE's own rule regimes ({n['b_nepse_lo']} "
         f"to {n['b_nepse_hi']}) and across the frontier markets examined ({n['b_lo']} to {n['b_hi']}). "
         "Checking it is not the same as plugging the estimate in: where it was well below one, setting it to "
         "zero did better than using the measured value.")
    edit(p, "Once the open is discounted by its measured reliability and the level is set on the close-to-close "
            "scale, the range can add information that close-to-close alone does not have: no classical range "
            "estimator beat the estimator anywhere, and on the indices, in Morocco and in the long Dhaka history "
            "it beat close-to-close too.",
         "Once the open is discounted and the level is set on the close-to-close scale, the range can add "
         "information that close-to-close alone does not have: under the plans' loss no classical range "
         "estimator had significantly lower loss than the estimator anywhere, and it beat close-to-close in "
         "the long Dhaka history, on the indices, as the classical range estimators also did there, and, less "
         "securely, in Morocco.")
    edit(p, "That the open-free form had the lowest loss of all nine estimators in Bangladesh, Vietnam and "
            "Morocco suggests something simpler: where the open overreacts,",
         "That the open-free form had the lowest five-session loss of all nine estimators in Bangladesh, "
         "Vietnam and Morocco suggests something simpler: where the open is unreliable,")
    edit(p, "Close-to-close was not beaten at short horizons in three of the four frontier markets, and it beat "
            "both forms immediately after NEPSE's rule change, which is when a regulator or a risk manager most "
            "needs a reliable number.",
         "Close-to-close was not beaten at short horizons in the primary panels of three of the four frontier "
         "markets, it beat both forms immediately after NEPSE's rule change, which is when a regulator or a "
         "risk manager most needs a reliable number, and the comparison with it turns on the loss function.")
    print("  corrected the Discussion")

    # ── Section 9 (M-020, M-025) ─────────────────────────────────────────────────────────────
    p = para_with(doc, "The estimator of Section 6.8 was developed and validated under three further frozen plans")
    edit(p, "was developed and validated under three further frozen plans",
         "was developed and tested under three further frozen plans")
    edit(p, "and its decision rules are applied mechanically to decision ledgers.",
         "and its decision rules are applied mechanically to decision ledgers. The plans are timestamped only "
         "by that history, which the author controls: they were not lodged with an external registry, and the "
         "history was rewritten once to correct its authorship, with each original commit identifier mapped to "
         "its replacement in the audit register.")
    edit(p, "Corrections made after these results were seen are recorded in the audit register.",
         "Corrections made after these results were seen are recorded in the audit register. A post hoc "
         "recheck (script 44; Tables 37 and 38), run after every verdict was known, varies the inference, the "
         "loss function and the calibration and tests the data assumptions; it changes no verdict.")
    print("  corrected Section 9")

    # ── Limitations: a sixteenth ─────────────────────────────────────────────────────────────
    p = para_with(doc, "Fifteenth, the estimator of Section 6.8 was designed on NEPSE data")
    para_after(p, (
        "Sixteenth, the estimator's record against close-to-close is fragile in places, as a post hoc recheck "
        "shows (Tables 37 and 38). Under the plans' QLIKE loss the frozen verdicts survive longer Newey-West "
        "lags except in Morocco, where the full form's advantage rests on one band regime and does not survive "
        "a Holm correction across the three plans; under an MSE loss close-to-close beats both forms in "
        "Vietnam (Table 37). The estimator's level on the close-to-close scale comes from its calibration, which the "
        "classical estimators can be given too. Three data assumptions were checked rather than taken on "
        f"trust: the repaired Dhaka dates look like the dates stamped correctly ({n['rep_big']}% and "
        f"{n['unr_big']}% of returns beyond 5%, and unbiasedness coefficients of {n['rep_b']} and "
        f"{n['unr_b']}); Morocco's band schedule, taken from press reports, matches the largest daily moves "
        f"in each regime (99.9th percentiles of {n['ma_q']}); and the coefficient is similar across Vietnam's "
        f"three exchanges, inferred from each share's largest moves ({n['vn_b_lo']} to {n['vn_b_hi']}). The "
        "Dhaka panels use the exchange's official close, which differs from the last trade on "
        f"{n['dse_close_diff']}% of the stock-days checked."), "Normal")
    print("  added a sixteenth limitation")

    # ── Conclusion ───────────────────────────────────────────────────────────────────────────
    p = para_with(doc, "The third finding is constructive.")
    edit(p, "The third finding is constructive. An estimator that measures how far the open can be trusted, "
            f"reads it only that far and calibrates to close-to-close variance ({NAME}, Section 6.8) was "
            "beaten by no classical range estimator in any of seven out-of-sample tests across Nepal, "
            "Bangladesh, Vietnam, Morocco and two benchmark indices, and outside NEPSE's post-reform regime its "
            f"level stayed within {n['lev_dev']}% of close-to-close variance, where the classical formulas "
            "missed it by tens of percent in both directions. It did not beat close-to-close at short horizons "
            "in three of the four frontier markets, and its open-free form, which leaves the open out "
            "altogether, is the version the frontier-market evidence supports.",
         "The third finding is constructive, and narrower than it first looks. An estimator that measures how "
         f"far the open can be trusted, reads it only that far and calibrates to close-to-close variance ({NAME}, "
         "Section 6.8) was not beaten significantly by any classical range estimator in seven out-of-sample "
         "tests across Nepal, Bangladesh, Vietnam, Morocco and two benchmark indices, under the loss function "
         f"the plans fixed. Its level stayed within {n['lev_dev']}% of close-to-close variance outside NEPSE's "
         "post-reform regime, but that is the calibration's doing: the classical formulas, which miss by tens "
         "of percent uncalibrated, come within a few percent when given the same calibration. It did not beat "
         "close-to-close at short horizons in the primary panels of three of the four frontier markets, its "
         "record against close-to-close changes with the loss function, and of its two forms the open-free one, "
         "which leaves the open out altogether, is the one the frontier-market evidence favours.")
    print("  corrected the Conclusion")

    doc.save(a.out)
    verify(a.out, n)


def verify(path, n) -> None:
    doc = docx.Document(path)
    paras = list(doc.paragraphs)
    body = "\n".join(p.text for p in paras)
    low = body.lower()
    bad = {k: v for k, v in BANNED.items() if k in low}
    assert not bad, f"banned phrases present: {bad}"
    assert paras[0].text == TITLE and paras[1].text == SUBTITLE
    assert len(doc.tables) == 38, f"expected 38 tables, found {len(doc.tables)}"
    assert len(doc.inline_shapes) == 8, f"expected 8 figures, found {len(doc.inline_shapes)}"
    assert abstract_words(doc) <= 255
    findings = next(p.text for p in paras if p.text.startswith("Findings:"))
    assert "NIFTY" not in findings, "M-014: no index figure in the abstract's findings"
    # the overstatements this round removes must not survive anywhere in the text
    for gone in ("Every analysis follows a plan frozen", "first two rule regimes alone", "before any later "
                 "observation was read", "before any holdout observation was read", "pre-registered test",
                 "determines what daily bars measure", "validated out of sample", "developed and validated",
                 "the data say that trust should be measured", "where the open is close to efficient",
                 "No classical range estimator beats", "beaten by no classical range estimator",
                 "many bars have no range at all", "Where one daily estimator must serve",
                 "What is added is that the open's weight is measured", "pooled medians between",
                 "Nepal, Bangladesh or Vietnam"):
        assert gone not in body, f"overstatement survives round 16: {gone!r}"
    for needle in ("Table 37.", "Table 38.", "Sixteenth,", "the calibration's doing", "post hoc recheck",
                   "Hansen and Lunde also weight the overnight return", "not lodged with an external registry",
                   "so close-to-close remains the primary measure here", "the plan lists what was known",
                   f"within {n['lev_dev']}% of close-to-close", "the failed predictions are reported as failures"):
        assert needle in body, f"missing after round 16: {needle!r}"
    for k in ("cls_cal_max", "anam_lag_max", "cls_lag_max", "C_cls_lo", "C_cls_hi", "sp_b", "nifty_b", "MA_nw2",
              "MA_holm", "nifty_nono", "VN21_mse_lo", "VN21_mse_hi", "VN5_mse", "VN5_mse_of", "A2C5_mse", "A2C5_mse_of",
              "C5_mse", "MA_4pct", "MA_rest_lo", "MA_rest_hi", "ma_zero", "b_lo", "b_hi", "rep_big", "unr_big",
              "rep_b", "unr_b", "dse_close_diff", "vn_b_lo", "vn_b_hi"):
        assert str(n[k]) in body, f"interpolated figure missing: {k} = {n[k]}"
    # every reading of the recheck is labelled post hoc
    for sentence in re.split(r"(?<=[.!?])\s+", body):
        if "MSE" in sentence or "Holm" in sentence:
            assert "post hoc" in sentence.lower() or "Table 37" in sentence or "Tables 37 and 38" in sentence, sentence[:100]
    i = find(paras, "Several limitations")
    j = find(paras, "11. Conclusion", i)
    found = [p.text.split(",")[0] for p in paras[i + 1:j] if p.text.split(",")[0] in ORD16]
    assert found == ORD16[1:len(found) + 1] and found[-1] == "Sixteenth", f"limitation ordinals: {found}"
    print(f"verification: {len(doc.tables)} tables, {len(doc.inline_shapes)} figures, abstract "
          f"{abstract_words(doc)} words; no overstatement from the recheck survives; post hoc readings labelled")


if __name__ == "__main__":
    main()
