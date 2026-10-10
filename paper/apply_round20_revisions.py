"""Round 20: Anam II, the market-implied open (plan M22), and Proposition 8 of the theory supplement.

    python paper/apply_round20_revisions.py [--base BASE.docx] [--out OUT.docx]

Plan M22 (``M22_ANAM2_PLAN.md``) was frozen before Anam II was computed on any test row and before any return was
computed on the Pakistan data. ``scripts/52_m22_evaluation.py`` ran it (tables 138-141), and
``scripts/53_m22_freeze.py`` froze its weights for the prospective test (table 142). The results are in
``M22_ANAM2_RESULTS.md``. The theory supplement gained Proposition 8 (``scripts/54``; tables 143-144), and the
step-by-step verification of the proofs now covers it (``scripts/46``; table 124).

This round:

* adds Section 6.9 on Anam II and Table 40, the plan's three hypotheses by panel;
* adds Anam II to the abstract, the introduction, Section 7.1, the Discussion, Section 9 and the Conclusion;
* adds an eighteenth limitation;
* updates three things in Section 9: the number of propositions, the proof count and the scripts that check them;
* adds the Pakistan file to the data statements and the tenth limitation;
* adds one reference: Bollerslev, Hood, Huss and Pedersen (2018).

RULES ENFORCED HERE

* Every number is interpolated from a table, and every claim is asserted against the tables before it is written.
* The test spans that were known when Anam II was designed are called seen. Pakistan, the market used at no stage
  of its design, is named as such.
* The plan's reporting rule is followed. Anam II is described by the reading of P1. The mechanism is stated beside
  it: the gain comes mostly from the dynamics.
* A sentence that reports Holm's adjustment names the plan that fixed it.
* Edits replace text inside the single run that holds it, so no paragraph loses its formatting.
"""
from __future__ import annotations

import argparse
import pathlib
import re
import sys

import pandas as pd

try:
    import docx
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

ROOT = pathlib.Path(__file__).resolve().parents[1]
TAB = ROOT / "output" / "tables"
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT / "src"))
from apply_referee_revisions import find, new_table_after, para_after, set_text  # noqa: E402
from apply_round4_revisions import set_repeat_header_row  # noqa: E402
from apply_round14_revisions import BANNED, ORDINALS, para_with, spacer_after, table_of  # noqa: E402
from apply_round15_revisions import NAME, SUBTITLE, TITLE, abstract_words  # noqa: E402
from apply_round16_revisions import edit  # noqa: E402

DOCX = ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"
NAME2 = "Anam II"
ORD18 = ORDINALS + ["Sixteenth", "Seventeenth", "Eighteenth"]
PK = "Pakistan 2016-2026"
PANELS = ["NEPSE", "DSE 2023-2026", "DSE 2009-2021", "Vietnam 2007-2020", "Morocco 2012-2026", PK]
SEEN = PANELS[:5]
INDICES = ["NIFTY50", "SP500"]
WORDS = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six", 7: "seven", 8: "eight", 9: "nine",
         10: "ten", 11: "eleven", 12: "twelve"}
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]
MARGIN = 0.01          # plan M22: a difference is inside the practical margin if |d| + 1.645 se < 1% of the reference
Z90 = 1.645

# Checked by web search, because the session's network could not reach Crossref (as for M-012, M-019 and the round
# 19 references): the DOI 10.1093/rfs/hhy041 and volume 31(7) are confirmed. The version of record's first page
# reads 2729-2773, while some indexing records give 2730 (audit register M-037).
REFS = [
    "Bollerslev, T., Hood, B., Huss, J., & Pedersen, L. H. (2018). Risk everywhere: Modeling and managing "
    "volatility. Review of Financial Studies, 31(7), 2729-2773. https://doi.org/10.1093/rfs/hhy041",
]


def t2(v) -> str:
    return f"{float(v):+.2f}".replace("-", "−")


def span(market: str) -> str:
    return "A2+C" if market == "NEPSE" else "test half"


def long_date(iso: str) -> str:
    y, m, d = (int(x) for x in iso.split("-"))
    return f"{d} {MONTHS[m - 1]} {y}"


def load() -> dict:
    """Every figure this round quotes, read from the tables, with each claim asserted first."""
    n = {}
    cl = pd.read_csv(TAB / "table139_m22_claims.csv")

    def h(hyp, mk, w):
        g = cl[(cl.hypothesis == hyp) & (cl.market == mk) & (cl.window == w)]
        assert len(g) == 1 and g.span.iloc[0] == span(mk)
        return g.iloc[0]

    for hyp in ("H1", "H2", "H3"):
        g = cl[cl.hypothesis == hyp]
        assert sorted(g.market.unique()) == sorted(PANELS) and len(g) == 12
        assert (g.verdict == "worse").sum() == 0, f"{hyp}: the text says Anam II is worse nowhere"

    # H1, the primary hypothesis, and P1, the plan's reading of it
    h1 = cl[cl.hypothesis == "H1"]
    better5 = set(h1[(h1.window == 5) & (h1.verdict == "better")].market)
    assert better5 == {"DSE 2023-2026", "Vietnam 2007-2020", PK}
    assert set(h1[(h1.window == 21) & (h1.verdict == "better")].market) == {"DSE 2023-2026"}
    n["h1_better5"] = WORDS[len(better5)]
    n["h1_t_dse"], n["h1_t_vn"], n["h1_t_pk"] = (t2(h("H1", mk, 5).t) for mk in
                                                 ("DSE 2023-2026", "Vietnam 2007-2020", PK))
    pk21 = h("H1", PK, 21)
    assert pk21.verdict == "no detectable difference" and pk21.p_holm > 0.05
    n["h1_t_pk21"], n["h1_p_pk21"] = t2(pk21.t), f"{pk21.p_holm:.3f}"
    assert h("H1", "NEPSE", 5).verdict == h("H1", "NEPSE", 21).verdict == "no detectable difference"
    p1 = cl.loc[cl.hypothesis == "P1", "verdict"].iloc[0]
    assert p1.startswith("supported") and "3 of 6" in p1 and "worse in 0" in p1

    # H2, the market-implied open with the dynamics fixed, and P2, the plan's predictions for it
    h2 = cl[cl.hypothesis == "H2"]
    b2 = set(zip(h2[h2.verdict == "better"].market, h2[h2.verdict == "better"].window.astype(int)))
    assert b2 == {("DSE 2023-2026", 5), ("DSE 2023-2026", 21), ("DSE 2009-2021", 21)}
    assert h2[h2.market == PK].practically_equivalent.all(), "the text says Pakistan's H2 differences are inside it"
    for mk in ("Vietnam 2007-2020", "Morocco 2012-2026"):
        assert h("H2", mk, 5).verdict == "no detectable difference", "P2's failed predictions"
    p2 = cl.loc[cl.hypothesis == "P2", "verdict"].iloc[0]
    held, of = re.match(r"(\d+) of (\d+) as predicted", p2).groups()
    n["p2_held"], n["p2_of"] = WORDS[int(held)], WORDS[int(of)]
    assert (n["p2_held"], n["p2_of"]) == ("three", "five")

    # H3, against the best return-only forecast
    h3 = cl[cl.hypothesis == "H3"]
    b3_5 = set(h3[(h3.window == 5) & (h3.verdict == "better")].market)
    b3_21 = set(h3[(h3.window == 21) & (h3.verdict == "better")].market)
    assert b3_5 == set(PANELS) - {"NEPSE"}
    n["h3_5"], n["h3_21"] = WORDS[len(b3_5)], WORDS[len(b3_21)]
    assert n["h3_21"] == "three"

    # the comparison table: the ablation reading, NEPSE's post-reform regime and the indices
    c = pd.read_csv(TAB / "table138_m22_comparison.csv")

    def row(mk, sp, w, model):
        g = c[(c.market == mk) & (c.span == sp) & (c.window == w) & (c.model == model)]
        assert len(g) == 1, (mk, sp, w, model)
        return g.iloc[0]

    for w in (5, 21):
        a = row("DSE 2023-2026", "test half", w, NAME2)["d_vs_HAR-open-free"]
        m = row("DSE 2023-2026", "test half", w, "HAR MSO (continuous)")["d_vs_HAR-open-free"]
        assert round(a, 4) == round(m, 4), "the text says the market-implied open alone gives the whole gain there"
    fo5 = {mk: row(mk, span(mk), 5, "FHARL open-free")["d_vs_HAR-open-free"] for mk in PANELS}
    ms5 = {mk: row(mk, span(mk), 5, "HAR MSO (continuous)")["d_vs_HAR-open-free"] for mk in PANELS}
    assert sum(fo5[mk] < ms5[mk] for mk in PANELS) >= 4, "the dynamics give more than the open in most panels"
    nc = row("NEPSE", "C", 5, NAME2)
    assert nc["t_vs_HAR-open-free"] < -1.96
    n["nepse_c_t"] = t2(nc["t_vs_HAR-open-free"])
    for mk in INDICES:
        for w in (5, 21):
            a, ref = row(mk, "test half", w, NAME2), row(mk, "test half", w, "HAR-open-free")
            se = abs(a["d_vs_HAR-open-free"] / a["t_vs_HAR-open-free"])
            assert abs(a["d_vs_HAR-open-free"]) + Z90 * se < MARGIN * ref.QLIKE_normalized, (mk, w)

    # the model confidence sets
    mcs = pd.read_csv(TAB / "table140_m22_mcs.csv")
    a2 = mcs[mcs.model == NAME2]
    prim = a2[a2.apply(lambda r: r.span == span(r.market), axis=1)]
    out = set(zip(prim[~prim.in_mcs_90].market, prim[~prim.in_mcs_90].window.astype(int)))
    assert out == {("Vietnam 2007-2020", 5)}
    vn = mcs[(mcs.market == "Vietnam 2007-2020") & (mcs.window == 5) & mcs.in_mcs_90]
    assert set(vn.model) == {"1/2 Anam II + 1/2 GJR"}
    for w in (5, 21):
        dse = mcs[(mcs.market == "DSE 2023-2026") & (mcs.window == w) & mcs.in_mcs_90]
        assert set(dse.model) == {NAME2, "1/2 Anam II + 1/2 GJR"}
        assert not a2[(a2.market == "NEPSE") & (a2.span == "C") & (a2.window == w)].in_mcs_90.iloc[0]

    # the panels: stale opens, and the Pakistan sample
    pan = pd.read_csv(TAB / "table141_m22_panels.csv").set_index("sample")
    n["dse_stale"] = f"{100 * pan.loc['DSE 2023-2026', 'stale_open_share']:.0f}"
    n["dse_oneprice"] = f"{100 * pan.loc['DSE 2023-2026', 'one_price_share']:.0f}"
    rest = pan.loc[[mk for mk in PANELS if mk != "DSE 2023-2026"], "stale_open_share"]
    assert pan.loc["DSE 2023-2026", "stale_open_share"] > rest.max()
    n["stale_lo"], n["stale_hi"] = f"{100 * rest.min():.0f}", f"{100 * rest.max():.0f}"
    pk = pan.loc[PK]
    assert int(pk["securities"]) == int(pk["rule: securities"]) == 101
    n["pk_secs"] = f"{int(pk['securities'])}"
    n["pk_days"] = f"{int(pk['rule: stock-days']):,}"
    n["pk_test_days"] = f"{int(pk['rule: test stock-days']):,}"
    n["pk_first"], n["pk_last"] = pk["rule: first session"], pk["rule: last session"]
    n["pk_test_from"] = long_date(pk["rule: first test session"])
    assert n["pk_first"].startswith("2016-10") and n["pk_last"].startswith("2026-10")

    # the open's reliability by component, on the training spans of the five seen frontier panels
    th = pd.read_csv(TAB / "table144_theory_market_open_data.csv").set_index("sample").loc[SEEN]
    assert th.seen.all() and (th.identity_gap.abs() < 1e-12).all()
    n["bM_lo"], n["bM_hi"] = f"{th.b_M.min():.2f}", f"{th.b_M.max():.2f}"
    n["bI_lo"], n["bI_hi"] = f"{th.b_I.min():.2f}", f"{th.b_I.max():.2f}"
    n["mu_lo"], n["mu_hi"] = f"{100 * th.mu.min():.0f}", f"{100 * th.mu.max():.0f}"
    assert (th.b_I < th.b_M).all() and th.b_M.min() > 0.5 and th.mu.max() < 0.25, "the text's three statements"
    assert ((th.b - th.b_I).abs() < (th.b - th.b_M).abs()).all(), "the pooled b is closer to b_I everywhere"

    # the development record, the theory checks and the proof verification
    led = pd.read_csv(ROOT / "output" / "dev_anam2" / "ledger.csv")
    n["variants"] = f"{led.groupby(['round', 'model']).ngroups}"
    ch = pd.read_csv(TAB / "table143_theory_market_open_checks.csv")
    assert bool(ch["pass"].all())
    n["m8_checks"] = f"{len(ch)}"
    steps = pd.read_csv(TAB / "table124_theory_proof_steps.csv")
    assert bool(steps["pass"].all()) and "8" in set(steps["proposition"].astype(str).str[0])
    n["steps"] = f"{len(steps)}"
    n["symbolic"] = f"{int(steps['method'].isin(['symbolic', 'symbolic-50dp']).sum())}"

    # the frozen weights, and the module constants the text states
    fw = pd.read_csv(TAB / "table142_m22_frozen_weights.csv")
    assert set(fw[fw.model == NAME2]["sample"]) == set(PANELS) | set(INDICES)
    from nepsevol import frontier as F
    from nepsevol.estimators import anam2 as A2
    from nepsevol import forecast_baselines as FB
    assert (A2.LAMBDA0, A2.FLOOR, FB.LONGRUN_SESSIONS) == (0.2, 0.05, 250)
    assert F.PSX_BANDS == ((None, 0.05), ("2020-01-20", 0.075), ("2024-05-27", 0.10)) and F.PSX_PRICE_FLOOR == 1.0
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default=str(DOCX))
    ap.add_argument("--out", default=str(DOCX))
    a = ap.parse_args()
    n = load()
    doc = docx.Document(a.base)
    if any(p.text.startswith("Table 40.") for p in doc.paragraphs):
        sys.exit("round 20 has already been applied to this document")
    if not any(p.text.startswith("Table 39.") for p in doc.paragraphs):
        sys.exit("round 19 must be applied first")

    p = para_with(doc, "Data and reproducibility note.")
    edit(p, "through apply_round19_revisions.py)", "through apply_round20_revisions.py)")
    edit(p, "and the forecast evaluations of Section 6.8 (plans M16-M18 and the corrected plan M20) follow",
         "and the forecast evaluations of Sections 6.8 and 6.9 (plans M16-M18, the corrected plan M20 and plan M22) "
         "follow")

    # ── abstract ─────────────────────────────────────────────────────────────────────────────
    p = para_with(doc, "Purpose:")
    edit(p, "Nepal has no exchange-traded options and no public intraday data, so volatility must be measured from "
            "the daily open-high-low-close bar.",
         "Nepal has no exchange-traded options or public intraday data, so volatility must be measured from daily "
         "open-high-low-close bars.")
    p = para_with(doc, "Design/methodology/approach:")
    edit(p, "Instrumented calibration slopes test for offsetting distortions, and an event study examines the 2026 "
            "rule package that widened NEPSE's pre-open band from ±2% to ±5%. A new estimator, designed on NEPSE's "
            "first two rule regimes, is tested on later regimes, two indices and stocks from Bangladesh, Vietnam and "
            "Morocco, and retested against return-only forecasts.",
         "Instrumented calibration slopes test for offsetting distortions; an event study examines the 2026 rule "
         "package that widened NEPSE's pre-open band from ±2% to ±5%. Two new estimators are tested out of sample "
         "in up to five frontier markets and two indices.")
    p = para_with(doc, "Findings:")
    edit(p, "the session undoes 64-87% of the overnight move, and after the rule package the share surviving fell by",
         "sessions undo 64-87% of the overnight move, and after the rule package the surviving share fell by")
    edit(p, "Anam's estimator weights the overnight move by the open's measured reliability and calibrates to "
            "close-to-close variance. After correction of the evaluation, its open-free form beats close-to-close at "
            "five sessions in six of seven samples, but the best range-based forecast beats the best return-only one "
            "in only two, and comparisons with close-to-close depend partly on the loss function.",
         f"{NAME}, in its open-free form, beats close-to-close at five sessions in six of seven samples, but "
         "return-only forecasts match or beat the best range-based forecast in five, and comparisons with "
         f"close-to-close depend partly on the loss function. {NAME2}, reading stale opens from other stocks' "
         f"overnight moves, beats the best earlier forecast at five sessions in {n['h1_better5']} of six frontier "
         f"panels, including unseen Pakistan, never loses, and beats the best return-only forecast in {n['h3_5']}.")
    edit(para_with(doc, "Originality/value:"),
         "an estimator built on that evidence is tested out of sample in four frontier markets.",
         "estimators built on that evidence are tested out of sample across five frontier markets.")
    print("  abstract")

    # ── introduction ─────────────────────────────────────────────────────────────────────────
    edit(para_with(doc, "This paper asks five questions"),
         "and retested against return-only forecasts after an independent audit corrected the forecast evaluation.",
         "and retested against return-only forecasts after an independent audit corrected the forecast evaluation; "
         "and a second estimator, tested under a further frozen plan in the same frontier markets and in a fifth, "
         "Pakistan, that played no part in its design.")
    edit(para_with(doc, "The fifth answer builds on the fourth."),
         "the best range-based forecast beats the best of them only in the long Dhaka history and in Morocco.",
         "the best range-based forecast beats the best of them only in the long Dhaka history and in Morocco. A second "
         f"estimator, {NAME2} (Section 6.9), follows from splitting the open into its market-wide and security-specific "
         "parts: the session keeps most of the first and less of the second. Where a security's open printed at the "
         "previous close, it reads the overnight move from the other securities' opens, and its forecast shares the "
         "market's volatility state across securities. Under a further plan frozen before it was tested, it has "
         "significantly lower loss than the corrected evaluation's most accurate forecast at five sessions in "
         f"{n['h1_better5']} of six frontier-market panels, including Pakistan's, a market that played no part in its "
         f"design, and significantly higher loss in none; it also beats the best return-only forecast at five sessions "
         f"in {n['h3_5']} of the six. Most of the gain comes from the dynamics: the market-implied open adds "
         "detectably only where stale prices are pervasive.")
    print("  introduction")

    # ── Section 6.9, after Section 6.8's last paragraph and before the table block ───────────
    last68 = para_with(doc, "Switching off one ingredient at a time locates what helps.")
    hd = para_after(last68, f"6.9 A second estimator: the market-implied open ({NAME2})", "Heading 2")
    new69 = [
        "Section 6.8 leaves two findings for a second estimator to build on: estimating the open's coefficient never "
        "beat setting it to zero, and the most accurate forecast in most samples paired the open-free kernel with HAR "
        "dynamics. Splitting each overnight return into the market's move, the date's cross-sectional mean, and the "
        "security's own remainder explains the first. On the training spans of the five frontier-market panels, the "
        f"session keeps most of the market part (its coefficient is between {n['bM_lo']} and {n['bM_hi']}) and, in "
        f"every panel, less of the security's own part (between {n['bI_lo']} and {n['bI_hi']}). But the market part is "
        f"only {n['mu_lo']}-{n['mu_hi']}% of the overnight second moment, so the pooled coefficient is close to the "
        "second: one coefficient shrinks the more reliable component by the less reliable one's factor. In a sample "
        "the split is an identity (theory supplement, Proposition 8). The market's move can be measured without a "
        "security's own opening error, from the other securities' opens.",
        f"The second estimator, which we call {NAME2}, uses that measurement where a security's own open says nothing "
        "about the overnight move: where it printed exactly at the previous close, as "
        f"{n['dse_stale']}% of opens did on the recent Dhaka panel's training span and between {n['stale_lo']}% and "
        f"{n['stale_hi']}% in the other panels. There the effective open o* is the equal-weighted mean overnight return "
        "of the panel's other securities on that date, which carries none of the security's own opening error; "
        "elsewhere o* = 0. The kernel is A = 0.8[o*² + R*²/D] + 0.2r², where R* = max(h, o*) - min(l, o*) is the range "
        "extended to reach the effective open, and D = 4 ln 2 except on a bar with a single price, where D = 1: a single "
        "observed increment has an expected square equal to its variance, and 4 ln 2 is the limit for a continuously "
        "observed range (Proposition 8). Where the open moved, A is the open-free kernel of Section 6.8 apart from that "
        "divisor. On a bar that never left the previous close the open-free kernel is zero, while A credits the "
        "market's move; this is the non-synchronous trading of Scholes and Williams (1977) and Lo and MacKinlay (1990), "
        f"seen in a daily bar. A is calibrated to close-to-close variance as {NAME} is.",
        f"{NAME2}'s forecast of the mean squared close-to-close return over the next h sessions is a factor HAR model of "
        "the calibrated kernel. Its components are the kernel's latest value and its means over the last 5 and 22 "
        "sessions; the security's own long-run mean of the kernel over 250 sessions; that long-run mean scaled by the "
        "date's cross-sectional median ratio of the 5-session, and of the 22-session, mean to the long-run mean, which "
        "carries the market's current volatility state to every security; and the security's own long-run mean of r². "
        "The weights are convex and fitted by minimising QLIKE on the training span, with at least 0.05 on the "
        "long-run mean so that every forecast is positive. The factor terms follow the common-component HAR models of "
        "Bollerslev, Hood, Huss and Pedersen (2018) and are not new; reading the market's move in place of a stale "
        f"opening print is what {NAME2} adds. A single series has no cross-section, so there {NAME2} is a HAR model "
        "of the open-free kernel with weights fitted by QLIKE.",
        f"{NAME2} was designed on the training spans of the seven samples alone, in a development record of "
        f"{n['variants']} variants that the package keeps, and plan M22 fixed it, its evaluation and its decision "
        "rules before it was computed on any test row (script 52). The test spans of Section 6.8 had nevertheless "
        "been seen: the corrected evaluation's results on them were known when it was designed. The plan therefore "
        "adds a market used at no stage of the design, the Pakistan Stock Exchange: "
        f"{n['pk_secs']} ordinary equities and {n['pk_days']} stock-days from October 2016 to October 2026, from a "
        "public compilation of the exchange's daily data pinned by digest, under panel rules fixed before any return "
        "was computed, including the exchange's dated price limits (5% or one rupee, whichever is larger, widened to "
        f"7.5% in January 2020 and to 10% in May 2024). Its test span starts on {n['pk_test_from']} and holds "
        f"{n['pk_test_days']} stock-days. Every comparison uses the corrected evaluation's targets, training rows, loss "
        "and standard errors on exactly its test rows, and plan M22 fixed Holm's adjustment across the six panels at "
        f"each horizon. H1, the primary hypothesis, compares {NAME2} with the corrected evaluation's most accurate "
        "forecast, the HAR model of the open-free kernel, rebuilt by the corrected evaluation's own code, which "
        "reproduced its reported test losses exactly before any comparison was made. H2 holds the dynamics fixed and replaces "
        f"the open-free kernel by {NAME2}'s, isolating the market-implied open. H3 compares {NAME2} with the best "
        "return-only forecast, chosen on the training span. A difference whose 90% interval lies within 1% of the "
        "reference forecast's loss is reported as inside a practical margin.",
        f"The plan's primary reading is supported (Table 40). At five sessions {NAME2} has significantly lower loss "
        f"than the HAR model of the open-free kernel in {n['h1_better5']} of the six panels, the recent Dhaka panel "
        f"(t = {n['h1_t_dse']}), Vietnam (t = {n['h1_t_vn']}) and Pakistan (t = {n['h1_t_pk']}), and significantly "
        "higher loss in none of the twelve panel-horizons. At 21 sessions it is better in the recent Dhaka panel only; "
        f"in Pakistan the difference (t = {n['h1_t_pk21']}) does not survive the adjustment (adjusted p = {n['h1_p_pk21']}). "
        f"Against the best return-only forecast it is better at five sessions in every panel except NEPSE, and at 21 "
        f"sessions in {n['h3_21']}. In NEPSE it shows no detectable difference from either on the primary span; after "
        f"the reform alone it beats the HAR model of the open-free kernel (t = {n['nepse_c_t']} at five sessions) but "
        "leaves the 90% model confidence set, where return-only forecasts remain, as in Section 6.8. Elsewhere it is "
        "in that set at every horizon except in Vietnam at five sessions, where only its equal-weighted combination "
        "with GJR-GARCH remains; in the recent Dhaka panel it and that combination are the only forecasts left. On the "
        f"two indices, where {NAME2} is a HAR model of the open-free kernel fitted by QLIKE, its differences from the "
        "corrected evaluation's version lie inside the practical margin.",
        "Most of the gain is the dynamics, not the market-implied open. With the dynamics held fixed (H2), the "
        "market-implied open lowers the loss detectably only where stale prices are pervasive: in the recent Dhaka "
        "panel at both horizons, where on the training span "
        f"{n['dse_stale']}% of opens printed at the previous close and {n['dse_oneprice']}% of bars had a single price, "
        f"and in the long Dhaka history at 21 sessions. {n['p2_held'].capitalize()} of the {n['p2_of']} predictions the "
        "plan made for H2 held; the predicted gains in Vietnam and Morocco were not detected, and in Pakistan the "
        "differences lie inside the practical margin. Read the other way, with the corrected evaluation's HAR dynamics "
        "fitted by QLIKE, the market-implied open alone gives the whole of "
        f"{NAME2}'s gain in the recent Dhaka panel (reported, no decision). In most other panels the factor HAR gives "
        "more of it. One reading, which the plan did not test, is that the factor terms lift the forecasts of "
        "securities whose own recent bars say little, which is the case the market-implied open was built for."
        f" {NAME2}'s weights, fitted on all the current data, are frozen for the "
        "prospective test of plan M21 (script 53).",
    ]
    anchor = hd
    for text in new69:
        anchor = para_after(anchor, text, "Normal")
    print("  Section 6.9")

    # ── Table 40, after Table 39 ─────────────────────────────────────────────────────────────
    doc.save(a.out)
    doc = docx.Document(a.out)
    anchor = spacer_after(table_of(doc, "Table 39."))
    t = pd.read_csv(TAB / "paper_table40_m22_anam2.csv", dtype=str).fillna("")
    cap = para_after(anchor, f"Table 40. {NAME2}, the market-implied open, under plan M22 (frozen before {NAME2} was "
                             "computed on any test row; script 52), on each panel's primary test span. Pakistan was "
                             "used at no stage of the design; the other panels' test spans had been seen. d is the mean "
                             "over stock-days of the difference in QLIKE loss (negative favours the first forecast) and t "
                             "its statistic with a standard error clustered by date; in plan M22, Holm's adjustment runs "
                             "across the six panels at each horizon, and 'inside the margin' marks a difference whose 90% "
                             "interval lies within 1% of the reference forecast's loss. H1 compares Anam II with the most "
                             "accurate forecast of the corrected evaluation, a HAR model of the open-free kernel; H2 holds "
                             "the dynamics fixed and compares the market-implied open with the open-free kernel; H3 "
                             "compares Anam II with the best return-only forecast chosen on the training span. The last "
                             "column gives membership of the 90% model confidence set.", "Normal")
    new_t = new_table_after(cap, doc, list(t.columns), t.values.tolist())
    set_repeat_header_row(new_t)
    spacer_after(new_t)
    print("  added Table 40")

    # ── Section 7.1 ──────────────────────────────────────────────────────────────────────────
    edit(para_with(doc, f"In NEPSE neither form of {NAME} (Section 6.8)"),
         "which the corrected evaluation found about as accurate in most samples.",
         f"which the corrected evaluation found about as accurate in most samples. In a panel of stocks, {NAME2} "
         f"(Section 6.9) improved on that HAR forecast at five sessions in {n['h1_better5']} of six frontier-market "
         "panels and was worse in none; in NEPSE the difference was not detectable.")
    print("  Section 7.1")

    # ── Discussion ───────────────────────────────────────────────────────────────────────────
    p = para_with(doc, "The estimator of Section 6.8 is the constructive side of these findings")
    edit(p, "and measure the range from the previous close. The limits are the benchmark's.",
         "and measure the range from the previous close. Section 6.9 qualifies this in one respect: the open's "
         "market-wide part is more reliable than its security-specific part, so the market's move can stand in "
         "for an open that never moved. That helps detectably where stale prices are pervasive; elsewhere most of "
         "the second estimator's gain comes from its dynamics, a factor HAR that carries the market's volatility state "
         "to every security. The limits are the benchmark's.")
    edit(p, "forecasts built from returns alone matched the range-based forecasts in most samples,",
         "forecasts built from returns alone matched the range-based forecasts of Section 6.8 in most samples, though "
         f"not {NAME2}'s at five sessions outside NEPSE,")
    print("  Discussion")

    # ── Section 9 ────────────────────────────────────────────────────────────────────────────
    p = para_with(doc, "The estimator of Section 6.8 was developed and tested under three further frozen plans")
    edit(p, "The statistics behind Sections 6.6-6.8 are set out", "The statistics behind Sections 6.6-6.9 are set out")
    edit(p, "seven propositions with proofs, derived after the plans above were run, the seventh after the audit "
            "described below.",
         "eight propositions with proofs, derived after the plans above were run, the seventh after the audit "
         "described below and the eighth after plan M22 was frozen.")
    old = re.search(r"\(\d+ steps, \d+ of them symbolic, all passing\)\.", p.text).group(0)
    edit(p, old, f"({n['steps']} steps, {n['symbolic']} of them symbolic, all passing).")
    edit(p, "Script 48 checks the seventh proposition, on how much of the opening error each kernel absorbs (63 checks, "
            "all passing).",
         "Script 48 checks the seventh proposition, on how much of the opening error each kernel absorbs (63 checks, "
         "all passing), and script 54 the eighth, on what the cross-section can recover of the overnight move "
         f"({n['m8_checks']} checks, all passing).")
    p = para_with(doc, "An independent audit of the package (9 October 2026) found four defects in the forecast evaluation")
    edit(p, "the Bangladeshi, Vietnamese and Moroccan results need third-party files",
         "the Bangladeshi, Vietnamese, Moroccan and Pakistani results need third-party files")
    para_after(
        p,
        f"{NAME2} (Section 6.9) was developed on the training spans alone, with every variant tried recorded "
        f"({n['variants']} in all, in the package's development record), and tested under plan M22 (scripts 52 and 53; "
        "tables 138-142). "
        "The plan was committed to the package's history together with the estimator's module, the evaluation script "
        "and the Pakistan reader before the evaluation was run on any test row; it records their digests, which the "
        "run checked. Two of its rules, the practical margin and the reading of a significant difference inside it as "
        "no difference, were added after a dry run on training rows had shown a negligible but significant "
        "difference; both are in the frozen plan. One development round, a boosted-tree forecast, did not finish and "
        "did not inform the design. The Pakistan file is a third-party compilation of the exchange's daily data that "
        f"the package pins by digest and does not redistribute. {NAME2}'s weights, fitted on all the current data, "
        "are frozen for the prospective test of plan M21 (table 142).",
        "Normal")
    print("  Section 9")

    # ── limitations ──────────────────────────────────────────────────────────────────────────
    edit(para_with(doc, "Tenth, on data provenance."),
         "the Bangladeshi, Vietnamese and Moroccan results of Section 6.8 need third-party files",
         "the Bangladeshi, Vietnamese, Moroccan and Pakistani results of Sections 6.8 and 6.9 need third-party files")
    p17 = para_with(doc, "Seventeenth, the forecast evaluation of the three plans")
    para_after(
        p17,
        f"Eighteenth, {NAME2} was designed when the corrected evaluation's results on its test spans were known, "
        "although it was never computed on them before its plan was frozen, so only its test in Pakistan, a market "
        "used at no stage of its design, and the prospective test are free of that knowledge. Its advantage comes "
        "mostly from dynamics whose factor terms follow existing work, and the market-implied open, the part that is "
        "new, helps detectably only where stale prices are pervasive. The Pakistan panel rests on a third-party "
        "compilation of the exchange's daily data, and the bars it places outside the exchange's limits, from "
        "unadjusted corporate actions or errors, are removed by rule rather than repaired.",
        "Normal")
    print("  limitations: tenth, and an eighteenth")

    # ── Conclusion ───────────────────────────────────────────────────────────────────────────
    edit(para_with(doc, "The third finding is constructive, and narrower than it first looks."),
         "is the one the frontier-market evidence favours.",
         f"is the one the frontier-market evidence favours. A second estimator, {NAME2}, reads a stale open from the "
         "other securities' opens and shares the market's volatility state across securities; under a further frozen "
         "plan it improves on the corrected evaluation's most accurate forecast at five sessions in "
         f"{n['h1_better5']} of six frontier-market panels, one of them Pakistan's, which played no part in its design, "
         f"and is worse in none, and it beats the best return-only forecast at five sessions in {n['h3_5']} of the six. "
         "Most of "
         "that gain is its dynamics rather than the market-implied open.")
    edit(para_with(doc, "The claims should stop there"),
         "and an estimator built on that evidence, tested out of sample under frozen plans and retested after an audit "
         "corrected their evaluation.",
         "and an estimator built on that evidence, tested out of sample under frozen plans, retested after an audit "
         "corrected their evaluation, and followed by a second that reads a stale open from the market's move, tested "
         "under a further frozen plan in a market that played no part in its design.")
    print("  Conclusion")

    # ── Declarations ─────────────────────────────────────────────────────────────────────────
    p = para_with(doc, "Data availability statement:")
    edit(p, "The analysis of Section 6.8 also uses the S&P 500 series", "The analyses of Sections 6.8 and 6.9 also use "
                                                                        "the S&P 500 series")
    edit(p, "stock-level files for the Dhaka, Vietnamese and Casablanca exchanges that are not redistributed",
         "stock-level files for the Dhaka, Vietnamese, Casablanca and Pakistan exchanges that are not redistributed")

    # ── References, kept alphabetical ────────────────────────────────────────────────────────
    paras = list(doc.paragraphs)
    first = find(paras, "Alizadeh, S., Brandt, M. W.")
    j = next(k for k, q in enumerate(paras) if q.text.startswith("Data and reproducibility note"))
    last = max(k for k in range(first, j) if paras[k].text.strip())
    existing = [paras[i].text.strip() for i in range(first, last + 1) if paras[i].text.strip()]
    assert not (set(existing) & set(REFS)), "a reference is already listed"
    for cited in ("Scholes, M., & Williams, J. (1977)", "Lo, A. W., & MacKinlay, A. C. (1990)"):
        assert any(r.startswith(cited) for r in existing), f"Section 6.9 cites {cited}, which must be listed"
    merged = sorted(set(existing) | set(REFS), key=lambda r: re.sub(r"^(The )", "", r).lower())
    slots = [i for i in range(first, last + 1) if paras[i].text.strip()]
    for kk, ref in enumerate(merged[:len(slots)]):
        set_text(paras[slots[kk]], ref)
    p_last = paras[slots[-1]]
    for ref in merged[len(slots):]:
        p_last = para_after(p_last, ref, "Normal")
    print(f"  reference list now {len(merged)} entries, alphabetical")

    doc.save(a.out)
    verify(a.out, n)


def verify(path, n) -> None:
    doc = docx.Document(path)
    paras = list(doc.paragraphs)
    body = "\n".join(p.text for p in paras)
    low = body.lower()
    bad = {k: v for k, v in BANNED.items() if k in low}
    assert not bad, f"banned phrases present: {bad}"
    assert paras[0].text == TITLE and paras[1].text == SUBTITLE
    assert len(doc.tables) == 40, f"expected 40 tables, found {len(doc.tables)}"
    assert len(doc.inline_shapes) == 8, f"expected 8 figures, found {len(doc.inline_shapes)}"
    words = abstract_words(doc)
    assert words <= 255, f"abstract has {words} words"
    findings = next(p.text for p in paras if p.text.startswith("Findings:"))
    for needle in ("Anam's estimator", "depend partly on the loss function", NAME2, "0.919 [0.760, 1.056]",
                   "fell by 0.333", "77 placebo dates"):
        assert needle in findings, f"the Findings must keep {needle!r}"
    assert "NIFTY" not in findings
    heads = [p.text for p in paras if p.style.name == "Heading 2"]
    i68 = heads.index(f"6.8 An estimator for an overreacting open: {NAME}")
    assert heads[i68 + 1] == f"6.9 A second estimator: the market-implied open ({NAME2})"
    for needle in ("Table 40.", "Eighteenth,", "plan M22", "apply_round20_revisions.py", "eight propositions with proofs",
                   f"({n['steps']} steps, {n['symbolic']} of them symbolic, all passing)",
                   f"({n['m8_checks']} checks, all passing)", "Bollerslev, Hood, Huss and Pedersen (2018)",
                   "Moroccan and Pakistani results", "Casablanca and Pakistan exchanges", f"t = {n['h1_t_pk']}",
                   "The plan's primary reading is supported (Table 40)"):
        assert needle in body, f"missing after round 20: {needle!r}"
    assert not re.search(r"ANAM|anam", body), "a file name would carry the eponym past the anonymiser"
    for gone in ("seven propositions with proofs", "(321 steps", "Bangladeshi, Vietnamese and Moroccan results",
                 "Dhaka, Vietnamese and Casablanca exchanges", "tested out of sample in four frontier markets."):
        assert gone not in body, f"superseded wording survives round 20: {gone!r}"
    for sentence in re.split(r"(?<=[.!?])\s+", body):
        if "MSE" in sentence or "Holm" in sentence:
            assert ("post hoc" in sentence.lower() or "Table 37" in sentence or "Tables 37 and 38" in sentence
                    or "Table 39" in sentence or "plan M20" in sentence or "Table 40" in sentence
                    or "plan M22" in sentence), sentence[:100]
    i = find(paras, "Several limitations")
    j = find(paras, "11. Conclusion", i)
    found = [p.text.split(",")[0] for p in paras[i + 1:j] if p.text.split(",")[0] in ORD18]
    assert found == ORD18[1:len(found) + 1] and found[-1] == "Eighteenth", f"limitation ordinals: {found}"
    refs_i = paras.index(next(p for p in paras if p.text == "References"))
    refs_j = next(k for k, p in enumerate(paras) if p.text.startswith("Data and reproducibility note"))
    refs = [p.text for p in paras[refs_i + 1:refs_j] if p.text.strip()]
    keys = [re.sub(r"^(The )", "", r).lower() for r in refs]
    assert keys == sorted(keys), "references must stay alphabetical"
    for r in REFS:
        assert r in refs, r[:40]
    t = pd.read_csv(TAB / "paper_table40_m22_anam2.csv", dtype=str).fillna("")
    cells = [[c.text for c in row.cells] for row in doc.tables[39].rows]
    assert cells[0] == list(t.columns) and cells[1:] == t.values.tolist(), "Table 40 must be the paper table"
    print(f"verification: {len(doc.tables)} tables, {len(doc.inline_shapes)} figures, abstract {words} words; "
          "Section 6.9 and Table 40 are in, Pakistan is named as the unseen market, the proof count is current")


if __name__ == "__main__":
    main()
