"""Two corrections closing M7: the premium claim, and the paired-comparison evidence.

    python paper/apply_round9_revisions.py [--base BASE.docx] [--out OUT.docx]

1. THE PREMIUM CLAIM WAS AN OVERREACH. Section 6.3 said the calibration slopes below one were
   "the familiar consequence of a volatility risk premium: India VIX sits above the realised
   volatility that follows it". A slope below one does not establish that. It says the RESPONSE
   of realised volatility to VIX is less than one-for-one -- attenuation, or regression toward
   the mean -- and is silent about the SIGN of the average forecast error. The direct evidence
   is the forecast error itself, which is now computed: mean and median VIX - RV with intervals,
   the mean VIX/RV ratio, the share of windows in which VIX exceeds RV, and the intercept and
   slope read jointly. That evidence does point the same way, but it is now reported as evidence
   and described as CONSISTENT WITH a volatility risk premium rather than as proof of one, with
   the competing explanations named.

2. THE COMPARISON IS PAIRED, AND THE EVIDENCE IS NOW SHOWN. The two correlations share their
   VIX observations, their windows, and outcomes that correlate at 0.889, so independent-sample
   inference would be invalid. The bootstrap always was paired -- one resampled index applied to
   all three series -- but the manuscript asserted the comparison rather than demonstrating it.
   The paired interval (width 0.234) is now shown beside the invalid independent one (0.648),
   together with a Fisher-z version of the same paired draw and Williams' t.

   Williams' t DISAGREES: at -2.587 it would reject equality in close-to-close's favour. It
   assumes trivariate normality, which strongly right-skewed volatility violates; both
   distribution-free paired bootstraps contain zero. The disagreement is reported rather than
   resolved by picking a side, and the frozen decision rule -- stated on the bootstrap interval
   -- still returns inconclusive. Every method agrees on one thing: nothing supports Parkinson
   over close-to-close.
"""

from __future__ import annotations

import argparse
import pathlib
import sys

import pandas as pd

try:
    import docx
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

ROOT = pathlib.Path(__file__).resolve().parents[1]
TAB = ROOT / "output" / "tables"
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from apply_referee_revisions import para_after, set_text, new_table_after  # noqa: E402
from apply_round4_revisions import set_repeat_header_row  # noqa: E402


def load_numbers():
    e = pd.read_csv(TAB / "table66_vix_forecast_error.csv").set_index("estimator")
    p = pd.read_csv(TAB / "table65_vix_paired_comparison.csv")

    def row(sub):
        return p[p.quantity.str.contains(sub, case=False, regex=False)].iloc[0]

    outcomes = row("correlation between the two OUTCOMES")
    pair = row("PAIRED bootstrap (used)")
    indep = row("INDEPENDENT bootstrap")
    z = row("Fisher z")
    will = row("Williams")
    cc, pk = e.loc["Close-to-close"], e.loc["Parkinson"]
    return {
        "cc_err": f"{cc.mean_VIX_minus_RV:.2f}",
        "cc_err_ci": f"[{cc.mean_lo95:.2f}, {cc.mean_hi95:.2f}]",
        "cc_ratio": f"{cc.mean_VIX_over_RV:.3f}",
        "cc_ratio_ci": f"[{cc.ratio_lo95:.3f}, {cc.ratio_hi95:.3f}]",
        "cc_share": f"{100 * cc.share_of_windows_VIX_above_RV:.0f}",
        "pk_err": f"{pk.mean_VIX_minus_RV:.2f}",
        "pk_err_ci": f"[{pk.mean_lo95:.2f}, {pk.mean_hi95:.2f}]",
        "pk_ratio": f"{pk.mean_VIX_over_RV:.3f}",
        "pk_ratio_ci": f"[{pk.ratio_lo95:.3f}, {pk.ratio_hi95:.3f}]",
        "pk_share": f"{100 * pk.share_of_windows_VIX_above_RV:.0f}",
        "cc_slope": f"{cc.slope:.3f}", "cc_int": f"{cc.intercept:.3f}",
        "pk_slope": f"{pk.slope:.3f}", "pk_int": f"{pk.intercept:.3f}",
        "r_out": f"{float(outcomes.value):.3f}",
        "pair_ci": f"[{pair.lo95:.3f}, {pair.hi95:.3f}]",
        "pair_w": f"{float(pair.hi95 - pair.lo95):.3f}",
        "indep_ci": f"[{indep.lo95:.3f}, {indep.hi95:.3f}]",
        "indep_w": f"{float(indep.hi95 - indep.lo95):.3f}",
        "z_ci": f"[{z.lo95:.3f}, {z.hi95:.3f}]",
        "will_t": f"{float(will.value):.2f}",
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default=str(ROOT / "paper" /
                    "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"))
    ap.add_argument("--out", default=str(ROOT / "paper" /
                    "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"))
    a = ap.parse_args()
    n = load_numbers()
    doc = docx.Document(a.base)

    # ── 1. replace the overreaching premium sentence ────────────────────────────────────────
    old = ("Calibration slopes are well below one (0.844 for close-to-close, 0.595 for "
           "Parkinson), the familiar consequence of a volatility risk premium: India VIX sits "
           "above the realised volatility that follows it.")
    new = ("Calibration slopes are well below one ({cc_slope} for close-to-close, {pk_slope} "
           "for Parkinson), which means realised volatility responds less than one-for-one to "
           "VIX. That on its own says nothing about whether VIX sits above or below what "
           "follows it, so the forecast error is reported directly. India VIX exceeds "
           "subsequent realised volatility by a mean of {cc_err} volatility points "
           "{cc_err_ci} against the close-to-close measure and {pk_err} {pk_err_ci} against "
           "Parkinson; the mean ratio of VIX to realised volatility is {cc_ratio} "
           "{cc_ratio_ci} and {pk_ratio} {pk_ratio_ci} respectively; and VIX is the larger of "
           "the two in {cc_share}% and {pk_share}% of windows. Read jointly, the intercepts "
           "({cc_int} and {pk_int}) and slopes place fitted realised volatility below VIX "
           "across the observed range. This is consistent with a volatility risk premium, but "
           "it does not prove one: differences in index construction, the two markets' trading "
           "calendars, the treatment of jumps, and the gap between a risk-neutral and a "
           "physical expectation can each push the same difference in the same direction. The "
           "size of the gap also depends on which realised measure is used — it is markedly "
           "larger against the range estimators than against close-to-close — which is itself a "
           "sign that construction, not only compensation for risk, is part of it "
           "(Table 26).").format(**n)
    hit = next((p for p in doc.paragraphs if old in p.text), None)
    if hit is not None:
        set_text(hit, hit.text.replace(old, new))
        print("  replaced the premium claim with forecast-error evidence")
    else:
        print("  !! premium sentence not found")

    # ── 2. add the paired-comparison paragraph after the ordering paragraph ─────────────────
    anchor = next((p for p in doc.paragraphs
                   if p.text.startswith("Neither estimator can be declared superior")), None)
    if anchor is not None and not any("must be a paired comparison" in p.text
                                      for p in doc.paragraphs):
        para_after(anchor,
            "That difference must be a paired comparison, and the inference behind it is "
            "reported rather than asserted (Table 27). The two correlations are computed from "
            "the same India VIX observations over the same windows, and their outcomes — "
            "forward Parkinson and forward close-to-close volatility — themselves correlate at "
            f"{n['r_out']}, so treating them as two independent correlations would be invalid. "
            "Every bootstrap replicate therefore draws one index and applies it to all three "
            "series together, keeping each window's triple intact. The consequence is large: "
            f"the paired interval on the difference is {n['pair_ci']}, of width {n['pair_w']}, "
            f"against {n['indep_ci']} and width {n['indep_w']} for the invalid independent "
            "calculation. A Fisher-transformed version of the same paired draw gives "
            f"{n['z_ci']}, also containing zero. One check disagrees: Williams' t for two "
            f"dependent correlations sharing a variable is {n['will_t']}, which would reject "
            "equality in close-to-close's favour. That test assumes trivariate normality, which "
            "strongly right-skewed volatility data violate, and both distribution-free paired "
            "intervals contain zero; the disagreement is reported rather than resolved by "
            "preference, and the decision rule fixed in advance was stated on the bootstrap "
            "interval. What no method supports, on any of these calculations, is Parkinson over "
            "close-to-close.", "Normal")
        print("  added the paired-comparison paragraph")

    doc.save(a.out)

    # ── Tables 26 and 27 ────────────────────────────────────────────────────────────────────
    doc = docx.Document(a.out)
    paras = list(doc.paragraphs)
    cap25 = next(p for p in paras if p.text.startswith("Table 25."))
    children = list(doc.element.body)
    pos = children.index(cap25._p)
    n_before = sum(1 for el in children[:pos] if el.tag.endswith("}tbl"))
    anchor_tbl = doc.tables[n_before]

    def spacer_after(tbl):
        el = OxmlElement("w:p")
        tbl._tbl.addnext(el)
        return Paragraph(el, tbl._parent)

    anch = spacer_after(anchor_tbl)
    for caption, csv_name in [
        ("Table 26. Forecast error: India VIX at t against realised volatility over "
         "(t, t+30 calendar days]. A calibration slope below one does not establish "
         "overprediction; these quantities do.",
         "paper_table26_vix_forecast_error.csv"),
        ("Table 27. Why the estimator comparison must be paired: the two correlations share "
         "their VIX observations and windows, and their outcomes are highly correlated.",
         "paper_table27_vix_paired_comparison.csv"),
    ]:
        t = pd.read_csv(TAB / csv_name, dtype=str).fillna("")
        cap = para_after(anch, caption, "Normal")
        new_t = new_table_after(cap, doc, list(t.columns), t.values.tolist())
        set_repeat_header_row(new_t)
        anch = spacer_after(new_t)
    print("  added Tables 26 and 27")

    doc.save(a.out)
    print(f"\nwrote {a.out}")

    # ── verification ────────────────────────────────────────────────────────────────────────
    doc = docx.Document(a.out)
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "the familiar consequence of a volatility risk premium" not in body, \
        "the overreaching premium claim survives"
    assert "consistent with a volatility risk premium, but it does not prove one" in body
    assert "must be a paired comparison" in body
    for t in ("Table 26.", "Table 27."):
        assert t in body, f"missing {t}"
    assert len(doc.tables) == 27, f"expected 27 tables, found {len(doc.tables)}"
    print(f"verification: premium claim qualified, paired inference documented "
          f"({len(doc.tables)} tables)")


if __name__ == "__main__":
    main()
