"""Round 17: the theory supplement, and the correction it forced in Appendix A.

    python paper/apply_round17_revisions.py [--base BASE.docx] [--out OUT.docx]

The statistics of the noisy open were derived in a theory supplement written in LaTeX
(``paper/theory/``). Each closed form was checked by ``scripts/45_theory_checks.py`` under plan
M19. Deriving them showed two things about Appendix A's bound on the opening error
(``AUDIT-REGISTER.md`` M-026):

  * the bound is valid but not sharp;
  * its statement that it "is attained when η is perfectly correlated with e_o" is wrong.

This round corrects that sentence. It states the sharp bound, (1 - b)^2 E[o^2], with its one-line
proof and its values (table122, post hoc). It also points Section 9 to the supplement and to the
outcome of plan M19's one test (table123).

RULES ENFORCED HERE
  * Every number is interpolated from a frozen table, and asserted before it is written.
  * The sharp bound and the theory are post hoc, and labelled so.
  * Edits replace text inside the single run that holds it.
"""
from __future__ import annotations

import argparse
import pathlib
import sys

import pandas as pd

try:
    import docx
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

ROOT = pathlib.Path(__file__).resolve().parents[1]
TAB = ROOT / "output" / "tables"
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from apply_round14_revisions import BANNED, para_with, replace_in  # noqa: E402
from apply_round15_revisions import SUBTITLE, TITLE, abstract_words  # noqa: E402
from apply_round16_revisions import WORDS, edit  # noqa: E402

DOCX = ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"
WRONG = "It is attained when η is perfectly correlated with e_o; under independence E[η²] = (1 - b)E[o²]. Both are reported"


def pct(v) -> str:
    return f"{100 * float(v):.1f}"


def load() -> dict:
    ap = pd.read_csv(TAB / "table122_theory_applications.csv")
    led = pd.read_csv(TAB / "table121_theory_checks.csv")
    pool = pd.read_csv(TAB / "table123_theory_pooling.csv")
    t97 = pd.read_csv(TAB / "table97_m15_posthoc.csv")

    def b1(sample, stat):
        r = ap[(ap.part == "B1") & (ap["sample"] == sample) & (ap.statistic == stat)].iloc[0]
        return r.value, r.lo, r.hi

    n = {}
    sharp, pub = "sharp bound (1-b)^2 E[o^2]/E[OC]", "published bound g(b) E[o^2]/E[OC]"
    for reg in ("A1", "C"):
        v, lo, hi = b1(f"NEPSE {reg}", sharp)
        n[f"{reg}_sharp"], n[f"{reg}_sharp_lo"], n[f"{reg}_sharp_hi"] = pct(v), pct(lo), pct(hi)
        # the published bound in table122 is table97's, recomputed on the same rows
        pv = b1(f"NEPSE {reg}", pub)[0]
        p97 = t97[(t97.regime == reg) & (t97.statistic == "lower bound on E[eta^2]/E[OC], any correlation with news")]["value"].iloc[0]
        assert abs(pv - p97) < 1e-9, (reg, pv, p97)
        assert v > pv, "the sharp bound must exceed the published one"
        n[f"{reg}_pub"] = pct(pv)
    assert len(led) >= 150 and bool(led["pass"].all()), "every check of the theory must pass before it is cited"
    n["checks"] = f"{len(led)}"
    v = pool["P4b_verdict"].value_counts()
    n["confirmed"] = WORDS.get(int(v.get("confirmed", 0)), str(int(v.get("confirmed", 0))))
    n["reversed"] = WORDS.get(int(v.get("reversed", 0)), str(int(v.get("reversed", 0))))
    n["cases"] = "ten" if len(pool) == 10 else str(len(pool))
    n["p4b"] = pool["P4b_overall"].iloc[0]
    assert n["p4b"] in ("supported", "not supported")
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default=str(DOCX))
    ap.add_argument("--out", default=str(DOCX))
    a = ap.parse_args()
    n = load()
    doc = docx.Document(a.base)
    if any("(1 - b)²E[o²] + E[(o* - b o)²]" in p.text for p in doc.paragraphs):
        sys.exit("round 17 has already been applied to this document")
    if not any(p.text.startswith("Table 37.") for p in doc.paragraphs):
        sys.exit("round 16 must be applied first")

    p = para_with(doc, "Data and reproducibility note.")
    replace_in(p, "through apply_round16_revisions.py)", "through apply_round17_revisions.py)")

    # ── Appendix A (M-026) ───────────────────────────────────────────────────────────────────
    p = para_with(doc, "Write the overnight and intraday returns as o = e_o + η - ε_prev")
    edit(p, WRONG,
         "It is not sharp, and earlier versions of this appendix were wrong to say that it is attained when η is "
         "perfectly correlated with e_o (corrected in round 17). Write o* = e_o - ε_prev for the efficient open "
         "measured from the printed previous close, so that o = o* + η and b o is the best linear predictor of o*. "
         "Then η = (1 - b)o - (o* - b o), and E[η²] = (1 - b)²E[o²] + E[(o* - b o)²] ≥ (1 - b)²E[o²]. This is the "
         "sharp bound. It is attained exactly when η is proportional to o*, and it exceeds the bound above for every "
         "b < 1. Under independence E[η²] = (1 - b)E[o²]. The weaker bound and the independence value are reported")
    edit(p, "where the bound was computed after the frozen results were seen.",
         "where the bound was computed after the frozen results were seen. The sharp bound, also post hoc, puts "
         f"the open's error at no less than {n['A1_sharp']}% [{n['A1_sharp_lo']}%, {n['A1_sharp_hi']}%] of the proxy "
         f"in the first regime, against {n['A1_pub']}% for the weaker bound, and at no less than {n['C_sharp']}% "
         f"[{n['C_sharp_lo']}%, {n['C_sharp_hi']}%] after the reform, against {n['C_pub']}% "
         "(table122_theory_applications.csv; proof in the theory supplement, Proposition 1).")
    print("  corrected Appendix A (M-026)")

    # ── Section 9: the theory supplement ────────────────────────────────────────────────────
    p = para_with(doc, "The estimator of Section 6.8 was developed and tested under three further frozen plans")
    tail = p.runs[-1].text.rstrip()
    assert tail.endswith("."), tail[-60:]
    p.runs[-1].text = (p.runs[-1].text.rstrip() +
                       " The statistics behind Sections 6.6-6.8 are set out in a theory supplement written in "
                       "LaTeX (paper/theory/): six propositions with proofs, derived after the plans above were "
                       f"run. Script 45 checked each closed form, identity and inequality against simulation or "
                       f"numerical integration ({n['checks']} checks, all passing). One prediction from the "
                       "supplement, that the open-free form gains most where a security's own open is noisiest, "
                       f"was fixed in advance under plan M19 and was {n['p4b']}: it was confirmed in "
                       f"{n['confirmed']} of {n['cases']} panel-horizon cases and reversed in {n['reversed']}.")
    print("  pointed Section 9 to the theory supplement")

    doc.save(a.out)
    verify(a.out, n)


def verify(path, n) -> None:
    doc = docx.Document(path)
    paras = list(doc.paragraphs)
    body = "\n".join(p.text for p in paras)
    low = body.lower()
    bad = {k: v for k, v in BANNED.items() if k in low}
    assert not bad, f"banned phrases present: {bad}"
    assert paras[0].text == TITLE and paras[1].text == SUBTITLE
    assert len(doc.tables) == 38, f"expected 38 tables, found {len(doc.tables)}"
    assert len(doc.inline_shapes) == 8, f"expected 8 figures, found {len(doc.inline_shapes)}"
    assert abstract_words(doc) <= 255
    assert "It is attained when η is perfectly correlated" not in body, "M-026: the wrong attainment claim survives"
    for needle in ("E[η²] = (1 - b)²E[o²] + E[(o* - b o)²] ≥ (1 - b)²E[o²]", "attained exactly when η is proportional to o*",
                   "apply_round17_revisions.py", "theory supplement written in LaTeX", "plan M19",
                   f"{n['A1_sharp']}% [{n['A1_sharp_lo']}%, {n['A1_sharp_hi']}%]", f"{n['C_sharp']}% [{n['C_sharp_lo']}%"):
        assert needle in body, f"missing after round 17: {needle!r}"
    print(f"verification: {len(doc.tables)} tables, {len(doc.inline_shapes)} figures; Appendix A corrected; "
          "Section 9 points to the theory supplement")


if __name__ == "__main__":
    main()
