"""Mandatory item 7: rebuild the India VIX section around a forward-looking test.

    python paper/apply_round8_revisions.py [--base BASE.docx] [--out OUT.docx]

The design and the decision rule were frozen in M7_ANALYSIS_PLAN.md before any forward result
was computed. Applying that rule mechanically to what came back:

    Parkinson - close-to-close correlation difference, primary specification:
        -0.064, 95% interval [-0.169, +0.065]  ->  contains zero  ->  INCONCLUSIVE

so the plan's second row applies: "Call the comparison inconclusive; no estimator-superiority
claim." Close-to-close is NOMINALLY ahead of Parkinson on every forward metric -- levels,
changes and the lead-lag profile -- and the manuscript now says so rather than leaving the
earlier impression that Parkinson led once an outlier was removed.

The finding that matters most is the lead-lag contrast. On the same origins, the same VIX
observation correlates far better with the 30 days BEFORE it than the 30 days AFTER it
(close-to-close 0.806 vs 0.694; Parkinson 0.778 vs 0.629). The trailing comparison the
manuscript previously reported was therefore measuring persistence and contemporaneous
co-movement, exactly as the review argued. It is retained, relabelled, and no longer carries any
forecasting or validation weight.
"""

from __future__ import annotations

import argparse
import pathlib
import sys

import pandas as pd

try:
    import docx
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

ROOT = pathlib.Path(__file__).resolve().parents[1]
TAB = ROOT / "output" / "tables"
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from apply_referee_revisions import find, para_after, set_text, new_table_after  # noqa: E402
from apply_round4_revisions import set_repeat_header_row  # noqa: E402


def load_numbers():
    f = pd.read_csv(TAB / "table59_vix_forward_primary.csv").set_index("estimator")
    c = pd.read_csv(TAB / "table60_vix_forward_changes.csv").set_index("estimator")
    ll = pd.read_csv(TAB / "table61_vix_leadlag.csv")
    ov = pd.read_csv(TAB / "table62_vix_forward_overlapping.csv").set_index("estimator")
    cmp_ = pd.read_csv(TAB / "table63_vix_parkinson_vs_cc.csv").iloc[0]

    def ll_get(est, which):
        r = ll[(ll.estimator == est) & ll.window.str.startswith(which)].iloc[0]
        return f"{float(r.pearson):.3f}"

    n = {
        "n_fwd": f"{int(f.loc['Parkinson', 'n_obs']):,}",
        "first": str(f.loc["Parkinson", "first_origin"]),
        "last": str(f.loc["Parkinson", "last_origin"]),
        "pk_fwd": f"{float(f.loc['Parkinson', 'pearson']):.3f}",
        "pk_fwd_ci": f"[{f.loc['Parkinson', 'pearson_lo95']:.3f}, "
                     f"{f.loc['Parkinson', 'pearson_hi95']:.3f}]",
        "cc_fwd": f"{float(f.loc['Close-to-close', 'pearson']):.3f}",
        "cc_fwd_ci": f"[{f.loc['Close-to-close', 'pearson_lo95']:.3f}, "
                     f"{f.loc['Close-to-close', 'pearson_hi95']:.3f}]",
        "pk_slope": f"{float(f.loc['Parkinson', 'slope_on_vix']):.3f}",
        "cc_slope": f"{float(f.loc['Close-to-close', 'slope_on_vix']):.3f}",
        "pk_chg": f"{float(c.loc['Parkinson', 'pearson_changes']):.3f}",
        "cc_chg": f"{float(c.loc['Close-to-close', 'pearson_changes']):.3f}",
        "pk_bwd": ll_get("Parkinson", "BACKWARD"),
        "pk_fwd_ll": ll_get("Parkinson", "FORWARD"),
        "cc_bwd": ll_get("Close-to-close", "BACKWARD"),
        "cc_fwd_ll": ll_get("Close-to-close", "FORWARD"),
        "diff": f"{float(cmp_.difference_pk_minus_cc):+.3f}",
        "diff_ci": f"[{cmp_.difference_lo95:.3f}, {cmp_.difference_hi95:.3f}]",
        "n_ov": f"{int(ov.loc['Parkinson', 'n_obs']):,}",
        "pk_ov": f"{float(ov.loc['Parkinson', 'pearson']):.3f}",
        "cc_ov": f"{float(ov.loc['Close-to-close', 'pearson']):.3f}",
    }
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default=str(ROOT / "paper" /
                    "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"))
    ap.add_argument("--out", default=str(ROOT / "paper" /
                    "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"))
    a = ap.parse_args()
    n = load_numbers()
    doc = docx.Document(a.base)

    # ── §6.3 heading ────────────────────────────────────────────────────────────────────────
    for p in list(doc.paragraphs):
        if p.text.strip() == "6.3 An options-market co-movement check: India VIX":
            set_text(p, "6.3 India VIX: a forward-looking test, and a trailing co-movement check")
            print("  retitled §6.3")
            break

    # ── §6.3 first paragraph: lead with the forward test ────────────────────────────────────
    for p in list(doc.paragraphs):
        if p.text.startswith("The most relevant external check for the paper's motivating"):
            set_text(p,
                "The most relevant external check for the paper's motivating question comes from "
                "India, where the NIFTY 50 has both cash-market OHLC data and an options-based "
                "volatility index. India VIX is a forward-looking, risk-neutral expectation of "
                "volatility over the following 30 calendar days; the estimators in this paper "
                "are backward-looking statistical measures of variation that has already "
                "happened. Those are different objects pointing in opposite directions in time, "
                "and a comparison that ignores the difference cannot support a forecasting "
                "claim. We therefore run the test the horizons actually license. For each origin "
                "date t we take India VIX observed at t and measure realised NIFTY volatility "
                "over the window (t, t+30 calendar days], excluding session t itself because "
                "India VIX is disseminated from quotes during that session. Windows are "
                f"non-overlapping in the primary specification, giving {n['n_fwd']} independent "
                f"forecast origins between {n['first']} and {n['last']}. Against forward realised "
                f"volatility, close-to-close correlates at {n['cc_fwd']} {n['cc_fwd_ci']} and "
                f"Parkinson at {n['pk_fwd']} {n['pk_fwd_ci']} (Table 23); in changes the "
                f"corresponding figures are {n['cc_chg']} and {n['pk_chg']}. Calibration slopes "
                f"are well below one ({n['cc_slope']} for close-to-close, {n['pk_slope']} for "
                "Parkinson), the familiar consequence of a volatility risk premium: India VIX "
                "sits above the realised volatility that follows it. The design, the reporting "
                "set and the decision rule applied below were fixed before any of these numbers "
                "were computed.")
            print("  rewrote §6.3 paragraph 1 around the forward test")
            break

    # ── §6.3 second paragraph: the ordering, the lead-lag, the reconciliation ───────────────
    for p in list(doc.paragraphs):
        if p.text.startswith("The ordering between them, however, is not robust"):
            set_text(p,
                "Neither estimator can be declared superior on this evidence. The difference in "
                f"forward correlation, Parkinson minus close-to-close, is {n['diff']} with a 95% "
                f"bootstrap interval of {n['diff_ci']}, which contains zero; under the "
                "pre-committed decision rule that outcome is inconclusive and no "
                "estimator-superiority claim follows. What can be said is that close-to-close is "
                "nominally ahead of Parkinson on every forward metric — levels, changes and the "
                "lead-lag profile — so nothing here supports the range estimator over the simple "
                "one. The lead-lag comparison is the more informative result. On the same "
                f"origins, the same VIX observation correlates {n['cc_bwd']} with the 30 days "
                f"BEFORE it against {n['cc_fwd_ll']} with the 30 days after it for "
                f"close-to-close, and {n['pk_bwd']} against {n['pk_fwd_ll']} for Parkinson "
                "(Table 24). The trailing comparison reported in earlier versions of this paper "
                "was therefore measuring persistence and contemporaneous co-movement rather than "
                "forecasting skill, and it is retained below only on that basis. Two numbers "
                "from those earlier versions should also be reconciled explicitly, because both "
                "have circulated: the Parkinson–VIX correlations of 0.776 and 0.832 differ only "
                "by whether the 5 October 2012 NSE flash-crash session is in the sample, and the "
                "R² values of 0.602 and 0.692 are not independent quantities at all but the "
                "squares of those two correlations. All four are trailing-horizon numbers and "
                "none of them is forecasting evidence.")
            print("  rewrote §6.3 paragraph 2 (ordering, lead-lag, reconciliation)")
            break

    # ── retain the trailing comparison, relabelled ──────────────────────────────────────────
    #
    # The paragraph rewrites above replace the text that used to carry the trailing numbers, so
    # they must be put back deliberately rather than allowed to disappear. The frozen plan is
    # explicit that they are retained -- they have been circulated, and quietly deleting a
    # circulated number is the failure mode this package keeps an audit register for.
    anchor = next((p for p in doc.paragraphs
                   if p.text.startswith("Neither estimator can be declared superior")), None)
    if anchor is not None and not any("Descriptive co-movement, horizon-mismatched"
                                      in p.text for p in doc.paragraphs):
        va = pd.read_csv(TAB / "table15_vix_anchor.csv").set_index("estimator")
        sens = pd.read_csv(TAB / "table31_nifty_outlier_sensitivity.csv")
        per = pd.read_csv(TAB / "table50_vix_period_sensitivity.csv")
        full = per[per.window.str.contains("full", case=False)].iloc[0]
        ovl = per[per.window.str.contains("overlap", case=False)].iloc[0]
        ex = sens[sens.iloc[:, 0].astype(str).str.contains("excluding")].iloc[0]
        para_after(anchor,
            "Descriptive co-movement, horizon-mismatched. The comparison reported in earlier "
            "versions of this paper is retained here, relabelled, because it has been "
            "circulated and should be reconcilable rather than quietly withdrawn. It correlates "
            "a TRAILING 21-session realised measure against a same-day India VIX that refers to "
            "the following 30 calendar days, so it is horizon-mismatched by construction, uses "
            "heavily overlapping windows, and is unsuitable for forecasting validation. On "
            f"{int(va.loc['Parkinson (21d)', 'n_obs']):,} matched observations over the full "
            f"India VIX history, the trailing Parkinson series correlates "
            f"{float(va.loc['Parkinson (21d)', 'corr']):.3f} with India VIX and the trailing "
            f"close-to-close series {float(va.loc['Close-to-close (21d)', 'corr']):.3f}. "
            f"Excluding the 5 October 2012 NSE flash-crash session — a genuine recorded session, "
            "retained in every reported result, whose extreme range enters the Parkinson series "
            "without a matching move in close-to-close volatility or the VIX — the Parkinson "
            f"figure becomes {float(ex.VIX_Parkinson_corr):.3f}. Restricting instead to the "
            f"{int(ovl.n_obs):,} observations from March 2024 onward that overlap the NEPSE "
            f"study period lowers both, to {float(ovl.VIX_Parkinson_corr):.3f} and "
            f"{float(ovl.VIX_CC_corr):.3f} (Table 13). These are descriptive statements about "
            "co-movement between a trailing measure and a forward-looking index. The forward "
            "test above is the one that speaks to forecasting, and it is the one the "
            "conclusions rest on.", "Normal")
        print("  restored the trailing comparison as a relabelled descriptive paragraph")

    # ── Abstract ────────────────────────────────────────────────────────────────────────────
    for p in list(doc.paragraphs):
        old = ("The India VIX co-movement check is real but period-sensitive (correlation "
               "0.78 over the full history versus 0.51 restricted to the study period) and "
               "establishes no ranking between estimators.")
        if old in p.text:
            set_text(p, p.text.replace(old,
                "Tested as an actual forecast — India VIX at t against realised NIFTY "
                "volatility over the following 30 days, non-overlapping windows — the "
                f"options-implied signal correlates {n['cc_fwd']} with subsequent close-to-close "
                f"and {n['pk_fwd']} with subsequent Parkinson volatility, well below the same "
                "measure's correlation with the 30 days preceding it; the estimator comparison "
                "is inconclusive and no superiority claim is made."))
            print("  updated the abstract")
            break

    # ── §10 limitation ──────────────────────────────────────────────────────────────────────
    for p in list(doc.paragraphs):
        if p.text.startswith("Fourth, the external India VIX exercise is cross-market"):
            set_text(p,
                "Fourth, the external India VIX exercise is evidence about the Indian market, "
                "not a Nepal-specific implied-volatility measurement, and it is reported at two "
                "horizons that must not be conflated. India VIX is a forward-looking, "
                "risk-neutral quantity; the estimators here are backward-looking statistical "
                "ones. The forward test of Section 6.3 matches those horizons and finds the "
                "estimator comparison inconclusive; the trailing 21-session comparison is "
                "horizon-mismatched, uses heavily overlapping windows, and is descriptive only. "
                "Neither establishes that a NEPSE volatility estimate would forecast NEPSE "
                "volatility, because Nepal has no options market against which such a test "
                "could be run — which is the paper's motivating problem, not a gap in its "
                "execution.")
            print("  updated the §10 VIX limitation")
            break

    # ── §8 Discussion ───────────────────────────────────────────────────────────────────────
    for p in list(doc.paragraphs):
        if p.text.startswith("The India VIX result is useful for interpretation, within limits."):
            set_text(p,
                "The India VIX result is useful for interpretation, within limits. Tested at the "
                "horizon India VIX actually refers to, the option-implied signal does carry "
                "information about the volatility that follows it, for both the range and the "
                "close-to-close measure, which helps explain why range-based measures remain "
                "informative in a market without options. But the forward correlations are "
                "materially below the trailing ones, the comparison between the two estimators "
                "is inconclusive, and close-to-close is nominally ahead throughout. The external "
                "evidence therefore establishes that daily OHLC data track an options-implied "
                "volatility state, and establishes no ranking between the estimators that track "
                "it. None of this makes a realized-volatility system a substitute for an options "
                "market: Nepal still lacks a forward-looking, risk-neutral volatility term "
                "structure, and a carefully constructed realized-volatility system substitutes "
                "for measurement, not for the information an option chain would provide.")
            print("  updated the §8 discussion paragraph")
            break

    # ── §11 conclusion ──────────────────────────────────────────────────────────────────────
    for p in list(doc.paragraphs):
        marker = "The NIFTY 50 and India VIX comparison further shows"
        if marker in p.text:
            i = p.text.find(marker)
            j = p.text.find("so that exercise establishes co-movement rather than dominance.")
            if i >= 0 and j > i:
                set_text(p, p.text[:i] + (
                    "The NIFTY 50 and India VIX comparison, tested at the 30-day horizon India "
                    "VIX actually refers to rather than against a trailing window, shows that "
                    "daily OHLC measures do track an options-implied volatility state into the "
                    f"future ({n['cc_fwd']} for close-to-close, {n['pk_fwd']} for Parkinson) "
                    "while tracking its immediate past more closely still, and it separates the "
                    "two estimators not at all: the difference between them is inconclusive, so "
                    "that exercise establishes co-movement rather than dominance."
                ) + p.text[j + len("so that exercise establishes co-movement rather than dominance."):])
                print("  updated the §11 conclusion")
            break

    doc.save(a.out)

    # ── Tables 23, 24, 25 ───────────────────────────────────────────────────────────────────
    doc = docx.Document(a.out)
    paras = list(doc.paragraphs)
    cap22 = next(p for p in paras if p.text.startswith("Table 22."))
    children = list(doc.element.body)
    pos = children.index(cap22._p)
    n_before = sum(1 for el in children[:pos] if el.tag.endswith("}tbl"))
    anchor_tbl = doc.tables[n_before]

    def spacer_after(tbl):
        el = OxmlElement("w:p")
        tbl._tbl.addnext(el)
        return Paragraph(el, tbl._parent)

    anchor = spacer_after(anchor_tbl)
    for caption, csv_name in [
        ("Table 23. India VIX as a forecast: correlation with realised NIFTY volatility over "
         "(t, t+30 calendar days], non-overlapping windows, session t excluded.",
         "paper_table23_vix_forward_primary.csv"),
        ("Table 24. Lead-lag: the same India VIX observation against the 30 days before it and "
         "the 30 days after it, on identical origins.",
         "paper_table24_vix_leadlag.csv"),
        ("Table 25. Forward-test sensitivity: daily overlapping windows with moving-block "
         "inference, and the direct Parkinson-versus-close-to-close comparison under the "
         "pre-committed decision rule.",
         "paper_table25_vix_forward_sensitivity.csv"),
    ]:
        t = pd.read_csv(TAB / csv_name, dtype=str).fillna("")
        cap = para_after(anchor, caption, "Normal")
        new_t = new_table_after(cap, doc, list(t.columns), t.values.tolist())
        set_repeat_header_row(new_t)
        anchor = spacer_after(new_t)
    print("  added Tables 23, 24, 25")

    doc.save(a.out)
    print(f"\nwrote {a.out}")

    # ── verification ────────────────────────────────────────────────────────────────────────
    doc = docx.Document(a.out)
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "forward-looking, risk-neutral" in body
    assert "horizon-mismatched" in body
    assert "inconclusive" in body
    for t in ("Table 23.", "Table 24.", "Table 25."):
        assert t in body, f"missing {t}"
    # the superiority claim must be gone
    for banned in ("reversing the ranking", "Parkinson's correlation rises to 0.832 while"):
        assert banned not in body, f"a superseded ordering claim survives: {banned!r}"
    assert len(doc.tables) == 25, f"expected 25 tables, found {len(doc.tables)}"
    print(f"verification: forward test, lead-lag and reconciliation in the text; "
          f"no ordering claim survives ({len(doc.tables)} tables)")


if __name__ == "__main__":
    main()
