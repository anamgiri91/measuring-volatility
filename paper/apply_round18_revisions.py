"""Round 18: the step-by-step verification of the theory, and the precision it asked for.

    python paper/apply_round18_revisions.py [--base BASE.docx] [--out OUT.docx]

``scripts/46_theory_proofs.py`` verified every part of every proposition in the theory supplement
(``paper/theory/``) and every step of its proofs. Where the step is algebra or calculus it did so
symbolically, with SymPy; elsewhere numerically, by simulation, or pathwise. The results are in
``output/tables/table124_theory_proof_steps.csv``.

It also covered the mathematical claims this paper makes itself. Reading the statements step by
step found one imprecision in the manuscript (``AUDIT-REGISTER.md`` M-027). Appendix A, as corrected
in round 17, says the sharp bound "is attained exactly when η is proportional to o*". That fails at
b = 0, where o* = 0 and the bound is attained with η = o. The exact condition is o* = b o, which for
b > 0 means η is proportional to o*.

This round:
  * states the exact condition in Appendix A;
  * adds the step-by-step verification to Section 9.

RULES ENFORCED HERE
  * Every number is interpolated from a table and asserted before it is written.
  * Edits replace text inside the single run that holds it.
"""
from __future__ import annotations

import argparse
import pathlib
import sys

import pandas as pd

try:
    import docx
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

ROOT = pathlib.Path(__file__).resolve().parents[1]
TAB = ROOT / "output" / "tables"
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from apply_round14_revisions import BANNED, para_with, replace_in  # noqa: E402
from apply_round15_revisions import SUBTITLE, TITLE, abstract_words  # noqa: E402
from apply_round16_revisions import edit  # noqa: E402

DOCX = ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx"
OLD = "It is attained exactly when η is proportional to o*, and it exceeds the bound above for every b < 1."
NEW = ("It is attained exactly when o* = b o, which for b > 0 means that η is proportional to o*, and it exceeds "
       "the bound above for every b < 1.")


def load() -> dict:
    steps = pd.read_csv(TAB / "table124_theory_proof_steps.csv")
    checks = pd.read_csv(TAB / "table121_theory_checks.csv")
    assert len(steps) >= 200 and bool(steps["pass"].all()), "every step must pass before it is cited"
    assert bool(checks["pass"].all())
    sym = int(steps["method"].isin(["symbolic", "symbolic-50dp"]).sum())
    return {"steps": f"{len(steps)}", "symbolic": f"{sym}", "checks": f"{len(checks)}"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default=str(DOCX))
    ap.add_argument("--out", default=str(DOCX))
    a = ap.parse_args()
    n = load()
    doc = docx.Document(a.base)
    if any(NEW in p.text for p in doc.paragraphs):
        sys.exit("round 18 has already been applied to this document")
    if not any("(1 - b)²E[o²] + E[(o* - b o)²]" in p.text for p in doc.paragraphs):
        sys.exit("round 17 must be applied first")

    p = para_with(doc, "Data and reproducibility note.")
    replace_in(p, "through apply_round17_revisions.py)", "through apply_round18_revisions.py)")

    edit(para_with(doc, "Write the overnight and intraday returns as o = e_o + η - ε_prev"), OLD, NEW)
    print("  stated the exact attainment condition in Appendix A (M-027)")

    old9 = f"numerical integration ({n['checks']} checks, all passing)."
    edit(para_with(doc, "The statistics behind Sections 6.6-6.8 are set out in a theory supplement"), old9,
         old9 + " Script 46 then verified each step of each proof, and the mathematical claims of this paper: "
         "the Yang-Zhang identity of Section 6.7, the non-negativity of Garman-Klass on a valid bar, the special "
         "cases of Anam's kernel and Appendix A. It did so symbolically wherever the step is algebra or calculus "
         f"({n['steps']} steps, {n['symbolic']} of them symbolic, all passing).")
    print("  added the step-by-step verification to Section 9")

    doc.save(a.out)
    verify(a.out, n)


def verify(path, n) -> None:
    doc = docx.Document(path)
    paras = list(doc.paragraphs)
    body = "\n".join(p.text for p in paras)
    bad = {k: v for k, v in BANNED.items() if k in body.lower()}
    assert not bad, f"banned phrases present: {bad}"
    assert paras[0].text == TITLE and paras[1].text == SUBTITLE
    assert len(doc.tables) == 38, f"expected 38 tables, found {len(doc.tables)}"
    assert len(doc.inline_shapes) == 8, f"expected 8 figures, found {len(doc.inline_shapes)}"
    assert abstract_words(doc) <= 255
    assert OLD not in body, "M-027: the imprecise attainment condition survives"
    for needle in (NEW, "apply_round18_revisions.py", f"({n['steps']} steps, {n['symbolic']} of them symbolic, all passing)"):
        assert needle in body, f"missing after round 18: {needle!r}"
    print(f"verification: {len(doc.tables)} tables, {len(doc.inline_shapes)} figures; Appendix A states the exact "
          "condition; Section 9 cites the step-by-step verification")


if __name__ == "__main__":
    main()
