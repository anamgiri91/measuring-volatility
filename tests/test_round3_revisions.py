"""Regression tests for the round-3 (second peer-review-evaluation) revisions.

Same standard as ``test_referee_revisions.py``: each test fails if a specific defect the
second-round critique identified -- and which was present in committed results -- comes back.
These target the SPECIFICATION of the new modules (equivalence margins, corporate-action
classification, horizon matching, block-date dependence, manifest honesty), not code coverage,
per SS21.
"""

from __future__ import annotations

import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))

import json
import re

import numpy as np
import pandas as pd
import pytest

from nepsevol.corporate_actions import (CLASSES, adjusted_previous_close,
                                        classify_disagreements, corporate_action_flags)
from nepsevol.equivalence import LOWER, MARGIN, UPPER, equivalence_verdict, verdict_sentence
from nepsevol.estimators import range_ as R
from nepsevol.inference import (ratio_of_sums_ci, ratio_of_sums_ci_block,
                                stationary_date_multiplicities)
from nepsevol.provenance import raw_data_status

ROOT = pathlib.Path(__file__).resolve().parents[1]
TAB = ROOT / "output" / "tables"


# ── item B: a STATED equivalence margin (post hoc), not "close to one" by eyeball ─────────

def test_margin_is_five_percent_and_symmetric():
    assert MARGIN == pytest.approx(0.05)
    assert LOWER == pytest.approx(0.95)
    assert UPPER == pytest.approx(1.05)


@pytest.mark.parametrize("lo,hi,want", [
    (0.96, 1.04, "equivalent"),
    (1.10, 1.30, "different"),      # lies entirely outside the margin
    (0.80, 1.30, "inconclusive"),   # straddles both margin boundaries: wide, not equivalence
    (0.94, 1.00, "inconclusive"),   # straddles the lower margin boundary
])
def test_equivalence_verdict_classifies_correctly(lo, hi, want):
    assert equivalence_verdict(lo, hi) == want


def test_a_wide_interval_covering_one_is_not_equivalence():
    """The exact failure mode the reviewer named: coverage of one is not evidence of parity.

    An interval that merely straddles one -- e.g. because it is wide and uninformative -- must
    not be reported as "equivalent". Only an interval that lies ENTIRELY inside the margin may,
    and a wide straddling interval is "inconclusive" (uninformative), not a positive finding of
    either equivalence or difference.
    """
    v = equivalence_verdict(0.5, 1.5)
    assert v != "equivalent"
    assert v == "inconclusive"


def test_verdict_sentence_never_says_unbiased_or_accurate():
    for lo, hi in [(0.96, 1.04), (1.10, 1.30), (0.94, 1.00)]:
        s = verdict_sentence("Parkinson", (lo + hi) / 2, lo, hi)
        for banned in ("unbiased", "accurate", "true bias"):
            assert banned not in s.lower()


def test_shipped_table7_uses_the_equivalence_verdict_not_ci_coverage():
    """Table 7's own reported ratios must be classified consistently with the margin, not by
    whether their two-way CI happens to contain one."""
    est = pd.read_csv(TAB / "table33_estimator_ratios_bootstrap.csv")
    for _, r in est.iterrows():
        assert r["equivalence_verdict"] == equivalence_verdict(r["lo95_twoway"], r["hi95_twoway"])


# ── item D: the 315 previous-close disagreements are classified, not merely counted ───────

def _panel_with_ca(n=10, seed=1):
    rng = np.random.default_rng(seed)
    dates = pd.date_range("2024-03-04", periods=n, freq="D")
    close = 100 * np.exp(np.cumsum(rng.normal(0, 0.01, n)))
    df = pd.DataFrame({"date": dates, "symbol": "AAA", "close": close,
                       "open": close, "high": close * 1.01, "low": close * 0.99,
                       "session_ord": np.arange(n)})
    df["prev_close"] = df["close"].shift(1)
    return df


def test_classes_are_exhaustive_and_mutually_exclusive():
    df = _panel_with_ca()
    # Row 5: engineer a genuine 1.20 bonus-style disagreement (published prev_close far below
    # the prior observed close), which must classify as corporate_action.
    df.loc[5, "prev_close"] = df.loc[4, "close"] / 1.20
    dis = classify_disagreements(df)
    assert len(dis) >= 1
    assert set(dis["ca_class"].unique()) <= set(CLASSES)


def test_corporate_action_requires_downward_revision_of_prev_close():
    """Only f > 1 + ROUNDING_TOL (published prev_close BELOW the prior close) may be classified
    corporate_action; an upward revision is a different, unexplained phenomenon."""
    df = _panel_with_ca()
    df.loc[5, "prev_close"] = df.loc[4, "close"] * 1.20   # ABOVE prior close
    dis = classify_disagreements(df)
    row = dis[dis.index == 5]
    assert not row.empty
    assert row["ca_class"].iloc[0] == "upward_adjustment"


def test_rounding_floor_is_not_swept_into_corporate_action():
    df = _panel_with_ca()
    df.loc[5, "prev_close"] = df.loc[4, "close"] * 1.001   # inside the 0.5% rounding floor
    dis = classify_disagreements(df)
    row = dis[dis.index == 5]
    assert not row.empty
    assert row["ca_class"].iloc[0] == "reference_rounding"


def test_adjusted_previous_close_only_moves_on_flagged_rows():
    df = _panel_with_ca()
    df.loc[5, "prev_close"] = df.loc[4, "close"] / 1.15
    unadj = adjusted_previous_close(df, use_corporate_actions=False)
    adj = adjusted_previous_close(df, use_corporate_actions=True)
    moved = (unadj != adj) & unadj.notna() & adj.notna()
    flags = corporate_action_flags(df)
    assert set(df.index[moved]) <= set(df.index[flags])
    assert moved.loc[5]


def test_adjusted_previous_close_is_nan_across_a_gap_even_with_corporate_actions():
    df = _panel_with_ca()
    df.loc[5, "session_ord"] = 10          # gap: sessions no longer advance by exactly 1
    df.loc[5, "prev_close"] = df.loc[4, "close"] / 1.15   # would look like a CA if not for the gap
    out = adjusted_previous_close(df, use_corporate_actions=True)
    assert pd.isna(out.loc[5])


def test_shipped_corporate_action_audit_accounts_for_all_315_disagreements():
    p = TAB / "table47_corporate_action_audit.csv"
    assert p.exists()
    audit = pd.read_csv(p).set_index("ca_class")
    assert int(audit["n_rows"].sum()) == 315
    assert set(audit.index) <= set(CLASSES)


# ── item C: Yang-Zhang's benchmark must share the estimator's own 21-session horizon ──────

def _yz_panel(n=60, seed=2):
    rng = np.random.default_rng(seed)
    o = 100 * np.exp(np.cumsum(rng.normal(0, 0.01, n)))
    c = o * np.exp(rng.normal(0, 0.01, n))
    h = np.maximum(o, c) * (1 + abs(rng.normal(0, 0.005, n)))
    l = np.minimum(o, c) * (1 - abs(rng.normal(0, 0.005, n)))
    return pd.DataFrame({"date": pd.date_range("2024-03-04", periods=n, freq="D"),
                         "symbol": "AAA", "open": o, "high": h, "low": l, "close": c,
                         "session_ord": np.arange(n)})


def test_yang_zhang_benchmark_is_rolling_21session_variance_not_same_session_return():
    df = _yz_panel()
    bench = R.yang_zhang_benchmark(df, window=21)
    same_session = R.close_to_close(df).pow(2)
    # The two series must differ (they answer different questions), and the benchmark must be
    # a smoother, rolling quantity: its non-missing count matches a 21-window rolling stat, not
    # a same-session return, which is missing for one fewer row.
    assert bench.notna().sum() < same_session.notna().sum()
    diffs = (bench - same_session).dropna()
    assert (diffs.abs() > 1e-12).any()


def test_horizon_matched_ratio_differs_from_same_session_ratio_on_shipped_data():
    """The reviewer's own recomputation: matching rows is not the same as matching horizons."""
    p = TAB / "table48_yang_zhang_horizon.csv"
    assert p.exists()
    t = pd.read_csv(p)
    same = t[(~t.horizon_matched) & (~t.adopted) &
             (t.previous_close_definition.str.startswith("unadjusted"))]
    horiz = t[(t.horizon_matched) & (~t.adopted) &
              (t.previous_close_definition.str.startswith("unadjusted"))]
    assert len(same) == 1 and len(horiz) == 1
    assert same.n_matched_rows.iloc[0] == horiz.n_matched_rows.iloc[0]     # rows ARE matched
    assert abs(same.sd_ratio.iloc[0] - horiz.sd_ratio.iloc[0]) > 0.005     # horizon still moves it


def test_exactly_one_adopted_row_in_the_horizon_table():
    t = pd.read_csv(TAB / "table48_yang_zhang_horizon.csv")
    assert t["adopted"].sum() == 1
    adopted = t[t["adopted"]].iloc[0]
    assert bool(adopted["horizon_matched"])
    assert adopted["previous_close_definition"].startswith("corporate-action")


# ── FOURTH-ROUND AUDIT: the previous close must reach the NUMERATOR, not only the benchmark ──
#
# The defect these guard against was real and shipped: `yang_zhang()` had no parameter through
# which the adopted corporate-action-adjusted previous close could reach its overnight term, so
# scripts/26_robustness.py built the numerator from the UNADJUSTED close while building the
# matched denominator from the ADJUSTED one, and the resulting mixed-definition ratio (1.309)
# was reported as though one definition applied throughout. The consistent specifications are
# 1.280 (adjusted) and 1.273 (unadjusted).

def test_yang_zhang_actually_uses_the_previous_close_it_is_given():
    """The parameter must change the answer, or it is decorative."""
    df = _panel_with_ca(n=40, seed=7)
    df["session_ord"] = range(len(df))
    ex = 20
    df.loc[ex, "prev_close"] = df.loc[ex - 1, "close"] / 1.20     # a bonus-style ex-date
    unadj = R.previous_session_close(df)
    adj = adjusted_previous_close(df)
    assert not adj.equals(unadj), "the fixture must actually contain a flagged ex-date"
    yz_unadj = R.yang_zhang(df, 21, prev_close=unadj)
    yz_adj = R.yang_zhang(df, 21, prev_close=adj)
    both = yz_unadj.notna() & yz_adj.notna()
    assert both.any()
    assert not np.allclose(yz_unadj[both], yz_adj[both]), \
        "yang_zhang ignored its prev_close argument -- the mixed-definition defect is back"


def test_yang_zhang_default_is_the_unadjusted_previous_session_close():
    """Omitting the argument must be the plain session close, so a caller that does not opt in
    to the adjustment cannot silently receive it (or vice versa)."""
    df = _panel_with_ca(n=40, seed=8)
    df["session_ord"] = range(len(df))
    df.loc[20, "prev_close"] = df.loc[19, "close"] / 1.15
    default = R.yang_zhang(df, 21)
    explicit = R.yang_zhang(df, 21, prev_close=R.previous_session_close(df))
    m = default.notna() & explicit.notna()
    assert m.any()
    assert np.allclose(default[m], explicit[m])


def test_shipped_yang_zhang_ratio_is_reproducible_from_one_stated_definition():
    """The adopted table value must equal a ratio computed with the SAME previous close on both
    sides -- which the 1.309 mixed-definition figure was not."""
    t = pd.read_csv(TAB / "table48_yang_zhang_horizon.csv")
    adopted = float(t[t["adopted"]].iloc[0]["sd_ratio"])
    est = pd.read_csv(TAB / "table33_estimator_ratios_bootstrap.csv").set_index("estimator")
    assert abs(float(est.loc["Yang-Zhang", "sd_ratio"]) - adopted) < 1e-9, \
        "Table 7's Yang-Zhang and the adopted cell of the horizon table disagree"
    # and it must not be the superseded mixed-definition value
    assert abs(adopted - 1.309) > 0.02, "the mixed-definition 1.309 figure is back"


# ── item G: the two-way bootstrap resamples dates i.i.d. and cannot see serial dependence ──

def test_stationary_date_multiplicities_sum_to_n_date():
    rng = np.random.default_rng(0)
    m = stationary_date_multiplicities(200, rng, mean_block=21)
    assert m.sum() == pytest.approx(200)
    assert (m >= 0).all()


def test_stationary_block_reduces_to_iid_at_mean_block_one():
    """A mean block length of 1 draws single dates -- the geometric distribution collapses to
    drawing one index at a time, so multiplicities look like an i.i.d. bootstrap's."""
    rng = np.random.default_rng(0)
    m = stationary_date_multiplicities(500, rng, mean_block=1)
    # With block length essentially always 1, this is indistinguishable in expectation from an
    # i.i.d. multinomial draw: no index should dominate the resample.
    assert m.max() < 15   # loose bound; a block bootstrap with long blocks would concentrate mass


def test_block_bootstrap_runs_and_brackets_the_point_estimate():
    n = 400
    rng = np.random.default_rng(3)
    sec = np.repeat(np.arange(20), n // 20)
    date = np.tile(np.arange(n // 20), 20)
    num = rng.normal(1.0, 0.1, n) ** 2
    den = rng.normal(1.0, 0.1, n) ** 2
    lo, hi = ratio_of_sums_ci_block(num, den, sec, date, n_boot=200, seed=1)
    point = np.sqrt(num.sum() / den.sum())
    assert lo <= point <= hi


def test_block_interval_is_not_narrower_than_the_iid_twoway_interval():
    """The whole point of the block bootstrap: absorbing serial dependence should not make the
    interval narrower than the (already dependence-aware) two-way i.i.d. one, on data with real
    within-block correlation."""
    n_dates, n_sec = 60, 15
    rng = np.random.default_rng(4)
    # Construct a market-wide shock that is correlated across ADJACENT dates (a "turbulent
    # week"), which only the block bootstrap can see.
    shock = np.zeros(n_dates)
    for i in range(0, n_dates, 5):
        shock[i:i + 5] = rng.normal(0, 0.3)
    date_idx = np.tile(np.arange(n_dates), n_sec)
    sec_idx = np.repeat(np.arange(n_sec), n_dates)
    eps = rng.normal(0, 0.05, n_dates * n_sec)
    r = 1.0 + shock[date_idx] + eps
    num, den = (r ** 2), np.ones_like(r)
    lo_two, hi_two = ratio_of_sums_ci(num, den, sec_idx, date_idx, n_boot=300, seed=5)
    lo_blk, hi_blk = ratio_of_sums_ci_block(num, den, sec_idx, date_idx, n_boot=300, seed=5,
                                            mean_block=5)
    assert (hi_blk - lo_blk) >= (hi_two - lo_two) - 1e-9


def test_shipped_block_interval_table_reports_all_six_estimators():
    t = pd.read_csv(TAB / "table49_interval_by_dependence.csv")
    assert set(t.estimator.unique()) == {"Parkinson", "Garman-Klass", "Rogers-Satchell",
                                         "AddRS", "Close-to-close", "Yang-Zhang"}
    assert set(t.clustering.unique()) >= {"security only (previous revision)"}


# ── item F: the build manifest must not claim a raw-data hash it does not have ────────────

def test_manifest_raw_data_claim_is_never_contradictory():
    raw = raw_data_status(ROOT)
    if raw.get("n_files") in (0, None):
        assert raw.get("aggregate_sha256") is None
        assert "claim" in raw and "NO claim" in raw["claim"]
    else:
        assert raw.get("aggregate_sha256") is not None


def test_shipped_manifest_matches_build_manifest_claim():
    p = ROOT / "data" / "processed" / "BUILD-MANIFEST.json"
    assert p.exists()
    shipped = json.loads(p.read_text())
    assert shipped["raw_data"]["n_files"] == 0
    assert shipped["raw_data"]["aggregate_sha256"] is None
    assert "NO claim" in shipped["raw_data"]["claim"]


# ── item A: the extreme thin tail must be visible, not smoothed into the quintile median ──

def test_shipped_thin_tail_table_separates_extreme_securities_from_the_rest():
    t = pd.read_csv(TAB / "table42_thin_tail.csv")
    groups = set(t.loc[t.n_securities.notna(), "group"])
    assert "extreme thin tail" in groups
    assert "all other ordinary equity" in groups
    extreme = t[t.group == "extreme thin tail"].iloc[0]
    other = t[t.group == "all other ordinary equity"].iloc[0]
    # The extreme tail must be a small population -- if it were most of the sample, carving it
    # out would not be a disclosure, it would be redefining the universe.
    assert extreme.n_securities < 10
    assert extreme.share_of_panel_pct < 5
    assert other.n_securities > extreme.n_securities


def test_named_extreme_securities_have_materially_higher_zero_range_than_the_aggregate():
    t = pd.read_csv(TAB / "table42_thin_tail.csv")
    named = t[t.participation.notna()]
    assert len(named) >= 3
    assert (named["zero_range_share"] > 0.15).all()


# ── item B (second half): information content is a different question from aggregate scale ─

def test_ratio_near_one_does_not_imply_high_correlation_with_the_proxy():
    """The reviewer's own point, verified against the shipped table: Parkinson and Rogers-
    Satchell can sit close to each other on SD ratio while tracking the proxy with very
    different fidelity."""
    t = pd.read_csv(TAB / "table43_information_content.csv").set_index("estimator")
    pk, rs = t.loc["Parkinson"], t.loc["Rogers-Satchell"]
    assert abs(pk["sd_ratio_to_proxy"] - rs["sd_ratio_to_proxy"]) < 0.05
    assert pk["pearson_vs_proxy"] - rs["pearson_vs_proxy"] > 0.3   # materially different tracking


# ── item E: the VIX co-movement figures must be reported over the NEPSE-overlap window too ─

def test_vix_period_sensitivity_reports_both_windows():
    t = pd.read_csv(TAB / "table50_vix_period_sensitivity.csv")
    windows = " ".join(t.window)
    assert "full" in windows.lower()
    assert "overlap" in windows.lower()
    overlap = t[t.window.str.contains("overlap", case=False)].iloc[0]
    full = t[t.window.str.contains("full", case=False)].iloc[0]
    # The overlap window is the actual NEPSE study period and must start no earlier than the
    # NEPSE panel itself (2024-03-04); the full window predates it by well over a decade.
    assert pd.Timestamp(overlap["first"]) >= pd.Timestamp("2024-03-04")
    assert pd.Timestamp(full["first"]) < pd.Timestamp("2020-01-01")


def test_overlap_correlations_are_materially_lower_than_full_sample():
    """Guards against the exact number the critique reported: correlations roughly halve once
    restricted to the period the NEPSE study actually covers."""
    t = pd.read_csv(TAB / "table50_vix_period_sensitivity.csv")
    overlap = t[t.window.str.contains("overlap", case=False)].iloc[0]
    full = t[t.window.str.contains("full", case=False)].iloc[0]
    assert overlap["VIX_Parkinson_corr"] < full["VIX_Parkinson_corr"] - 0.15
    assert overlap["VIX_CC_corr"] < full["VIX_CC_corr"] - 0.15


# ── FOURTH-ROUND MANDATORY ITEMS 3, 5 AND 11 ─────────────────────────────────────────────────

def test_margin_is_documented_as_post_hoc_not_prespecified():
    """M3: the margin is applied consistently but was declared after the estimates existed.
    Calling it 'prespecified' claims a protection against selection that does not exist here."""
    import nepsevol.equivalence as eq
    doc = (eq.__doc__ or "")
    assert "DECLARED POST HOC" in doc
    assert "NOT prespecified" in doc
    for s in (verdict_sentence("X", 1.0, 0.99, 1.01),
              verdict_sentence("X", 1.3, 1.2, 1.4),
              verdict_sentence("X", 0.99, 0.94, 1.00)):
        assert "prespecified" not in s


def test_margin_grid_is_reported_and_brackets_the_headline_margin():
    """M3: no verdict may rest on a single post hoc margin."""
    from nepsevol.equivalence import MARGIN, MARGINS
    assert MARGIN in MARGINS
    assert min(MARGINS) < MARGIN < max(MARGINS)
    t = pd.read_csv(TAB / "table54_equivalence_margin_sensitivity.csv")
    for m in MARGINS:
        assert f"verdict_at_{100 * m:g}pct" in t.columns
    # both confidence levels must be present, since 95% is conservative and 90% is TOST
    assert t.interval.str.contains("95%").any() and t.interval.str.contains("90%").any()


def test_a_verdict_that_flips_across_the_margin_grid_is_visible():
    """The grid is only worth reporting if it can disagree with itself -- and here it does."""
    t = pd.read_csv(TAB / "table54_equivalence_margin_sensitivity.csv")
    mcols = [c for c in t.columns if c.startswith("verdict_at_")]
    assert t[mcols].nunique(axis=1).gt(1).any(), \
        "no verdict changes across the margin grid: the sensitivity table would be vacuous"


def test_partial_year_regime_is_flagged_unusable_as_an_annualisation_factor():
    """M5: annualising 99 sessions observed over 0.39 of a year is an extrapolation, and both
    regimes schedule five sessions a week, so the reform cannot change the annual count."""
    t = pd.read_csv(TAB / "table46_annualisation_regimes.csv").set_index("regime")
    assert "usable_as_annualisation_factor" in t.columns
    monfri = t.loc["Mon-Fri (from 2026-04-06)"]
    assert not bool(monfri.usable_as_annualisation_factor)
    assert float(monfri.span_years) < 1.0
    assert bool(t.loc["WHOLE SAMPLE (the reported A)"].usable_as_annualisation_factor)


def test_annualisation_regime_labels_match_their_own_dates():
    """M5: the interpretation strings were once assigned POSITIONALLY to an alphabetically
    sorted frame, which swapped 'pre-reform' and 'post-reform' between the two regime rows."""
    t = pd.read_csv(TAB / "table46_annualisation_regimes.csv").set_index("regime")
    monfri = t.loc["Mon-Fri (from 2026-04-06)"]
    sunthu = t.loc["Sun-Thu (to 2026-04-05)"]
    assert "POST-reform" in monfri.interpretation and "Monday-Friday" in monfri.interpretation
    assert "PRE-reform" in sunthu.interpretation and "Sunday-Thursday" in sunthu.interpretation
    assert pd.Timestamp(monfri.first) > pd.Timestamp(sunthu.last)


def test_weighting_is_reported_for_every_named_estimator():
    """M11: equal-security aggregation existed only for Parkinson and RS, by quintile."""
    t = pd.read_csv(TAB / "table52_weighting_all_estimators.csv")
    assert set(t.estimator) == {"Parkinson", "Garman-Klass", "Rogers-Satchell",
                                "AddRS", "Close-to-close", "Yang-Zhang"}
    for col in ("stockday_weighted", "equal_security_mean", "equal_security_median",
                "security_p05", "security_p95"):
        assert t[col].notna().all(), f"{col} missing for some estimator"
    # the security-level band must be a real band, not a degenerate point, for the estimators
    # whose ratio is not one by construction
    non_trivial = t[t.estimator != "Close-to-close"]
    assert (non_trivial.security_p95 - non_trivial.security_p05 > 0.05).all()


def test_stockday_and_equal_security_weightings_agree_within_a_reported_bound():
    """M11: if the two weightings diverged materially the headline would be a weighting artifact.
    They do not -- and the bound is asserted here so a future change cannot pass silently."""
    t = pd.read_csv(TAB / "table52_weighting_all_estimators.csv")
    assert float(t.weighting_gap.max()) < 0.05


# ── MANDATORY ITEM 4: the thinness screen must not select on the estimator's own outcome ─────

def test_lagged_screen_reads_no_estimator_output():
    """The screen's source must not touch price, range, return or any estimator column beyond
    what it needs to REPORT the result. Selection reads participation and trade counts only."""
    src = (ROOT / "scripts" / "31_lagged_thinness_screen.py").read_text()
    sel = src.split("def build_flags")[1].split("def ratio_block")[0]
    for banned in ("v_pk", "v_oc", "zero_range", "high", "low", "close", "open"):
        assert banned not in sel, f"the selection rule reads {banned!r}: it is not outcome-independent"


def test_lagged_screen_history_floor_is_scheduled_not_traded_sessions():
    """A floor counted in TRADED sessions would condition eligibility on the activity being
    screened, reintroducing the circularity in a subtler form."""
    src = (ROOT / "scripts" / "31_lagged_thinness_screen.py").read_text()
    elig = src.split("def eligibility")[1].split("def build_flags")[0]
    assert "first_sord" in elig and "traded" not in elig


def test_lagged_screen_does_not_reproduce_the_zero_range_collapse():
    """The substantive finding: an outcome-independent screen does not find the collapse that
    the zero-range grouping reported, which is why the latter is descriptive only."""
    lagged = pd.read_csv(TAB / "table55_lagged_thinness_baseline.csv")
    thin = lagged.iloc[0]
    zr = pd.read_csv(TAB / "table42_thin_tail.csv")
    zr_thin = zr[zr.group == "extreme thin tail"].iloc[0]
    assert float(zr_thin.Parkinson) < 0.90, "the zero-range grouping no longer shows a collapse"
    assert float(thin.Parkinson) > 0.95, \
        "the lagged screen now shows a collapse; the manuscript's claim must be revisited"


def test_lagged_screen_sensitivity_grids_are_complete_and_stable():
    grid = pd.read_csv(TAB / "table56_lagged_thinness_grid.csv")
    assert set(grid.participation_threshold) == {0.80, 0.90, 0.95}
    assert set(grid.trade_count_tail) == {0.05, 0.10, 0.20}
    assert len(grid) == 9
    floors = pd.read_csv(TAB / "table57_lagged_thinness_history_floor.csv")
    assert set(floors.history_floor_scheduled_sessions) == {40, 60, 120}
    comps = pd.read_csv(TAB / "table58_lagged_thinness_components.csv")
    assert set(comps.component) == {"participation only", "trade count only",
                                    "composite OR (baseline)"}
    # no cell of the grid or the floor sweep may show a collapse
    assert (grid.Parkinson_thin > 0.95).all()
    assert (floors.Parkinson_thin > 0.95).all()


def test_manuscript_calls_the_screen_post_hoc_not_ex_ante():
    """'Ex ante' would be read as a claim that the rule predated the results. It did not."""
    import docx
    d = docx.Document(ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx")
    body = "\n".join(p.text for p in d.paragraphs)
    assert "lagged, outcome-independent liquidity screen" in body
    assert "declared post hoc and applied consistently" in body
    assert "ex ante" not in body.lower()
    assert "descriptive failure-case inventory" in body


# ── MANDATORY ITEM 7: the forward window must not see its own origin or its past ─────────────

def _load_nifty_for_forward():
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "vixfwd", ROOT / "scripts" / "32_vix_forward_validation.py")
    return spec


def test_forward_window_excludes_the_origin_session():
    """The single most important guard: the outcome window starts at origin+1. India VIX is
    disseminated from quotes during session t, so session t's own return is contemporaneous
    with the forecast's formation, not subsequent to it."""
    src = (ROOT / "scripts" / "32_vix_forward_validation.py").read_text()
    fn = src.split("def forward_rv")[1].split("def ")[0]
    assert "lo = origin_idx + 1" in fn, \
        "the forward window no longer starts strictly after its origin"
    assert "origin_idx - " not in fn and "shift(" not in fn, \
        "the forward window reaches backwards"


def test_forward_and_backward_windows_are_built_by_different_code_paths():
    """The backward window deliberately CONTAINS t; the forward one must not. If both were
    built by one function with a sign flip, a future edit could silently align them."""
    src = (ROOT / "scripts" / "32_vix_forward_validation.py").read_text()
    fwd = src.split("def forward_rv")[1].split("def ")[0]
    bwd = src.split("def backward_rv")[1].split("def ")[0]
    assert "lo = origin_idx + 1" in fwd
    assert "hi = origin_idx + 1" in bwd


def test_shipped_forward_correlations_are_lower_than_backward_ones():
    """The substantive finding of item 7: the trailing comparison was measuring persistence.
    If forward ever matched backward, the horizon-mismatch criticism would be moot -- and this
    test would need revisiting rather than silently passing."""
    ll = pd.read_csv(TAB / "table61_vix_leadlag.csv")
    for est in ll.estimator.unique():
        f = ll[(ll.estimator == est) & ll.window.str.startswith("FORWARD")].pearson.iloc[0]
        b = ll[(ll.estimator == est) & ll.window.str.startswith("BACKWARD")].pearson.iloc[0]
        assert f < b, f"{est}: forward correlation is not below backward"


def test_primary_forward_specification_is_non_overlapping():
    """Overlapping windows share up to 29 days of outcome data; the headline must not rest on
    an interval that assumes independence it does not have."""
    prim = pd.read_csv(TAB / "table59_vix_forward_primary.csv")
    ov = pd.read_csv(TAB / "table62_vix_forward_overlapping.csv")
    assert int(prim.n_obs.iloc[0]) < int(ov.n_obs.iloc[0]) / 10, \
        "the primary sample is not materially smaller than the overlapping one"


def test_vix_number_reconciliation_shows_R2_is_just_corr_squared():
    """0.692 and 0.602 are not an independent discrepancy: they are 0.832^2 and 0.776^2."""
    rec = pd.read_csv(TAB / "table64_vix_number_reconciliation.csv")
    assert rec.R2_equals_corr_squared.all()
    vals = set(rec.reported_value.round(3))
    assert {0.776, 0.832} <= vals


def test_manuscript_does_not_call_the_trailing_comparison_forecasting():
    import docx
    d = docx.Document(ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx")
    body = "\n".join(p.text for p in d.paragraphs)
    assert "horizon-mismatched" in body.lower() or "horizon mismatch" in body.lower()
    assert "forward-looking, risk-neutral" in body


# ── M7 close-out: the premium claim and the paired comparison ────────────────────────────────

def test_premium_is_evidenced_by_forecast_error_not_by_the_slope():
    """A calibration slope below one says the response is less than one-for-one. It is silent
    about the SIGN of the forecast error, so it cannot establish overprediction on its own."""
    e = pd.read_csv(TAB / "table66_vix_forecast_error.csv")
    for col in ("mean_VIX_minus_RV", "mean_lo95", "mean_hi95", "mean_VIX_over_RV",
                "ratio_lo95", "ratio_hi95", "share_of_windows_VIX_above_RV",
                "intercept", "slope"):
        assert col in e.columns, f"the forecast-error evidence lacks {col}"
    assert (e.mean_lo95 < e.mean_VIX_minus_RV).all()
    assert (e.mean_VIX_minus_RV < e.mean_hi95).all()


def test_manuscript_does_not_claim_the_slope_proves_a_premium():
    import docx
    d = docx.Document(ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx")
    body = "\n".join(p.text for p in d.paragraphs)
    assert "the familiar consequence of a volatility risk premium" not in body
    assert "consistent with a volatility risk premium, but it does not prove one" in body
    # the competing explanations must be named, not just gestured at
    for alt in ("index construction", "trading", "jumps", "risk-neutral"):
        assert alt in body, f"the premium caveat does not name {alt!r}"


def test_estimator_comparison_bootstrap_is_paired():
    """Structural: one resampled index must be applied to VIX and BOTH outcomes together.
    Two independent indices would ignore that the correlations share observations."""
    src = (ROOT / "scripts" / "32_vix_forward_validation.py").read_text()
    blk = src.split("# ── THE DIRECT COMPARISON")[1].split("# ── the comparison must be PAIRED")[0]
    assert "idx = rng.integers(0, n, n)" in blk
    assert "xv[idx], yp[idx]" in blk and "xv[idx], yc[idx]" in blk, \
        "the two correlations are no longer resampled with the same index"


def test_paired_interval_is_narrower_than_the_invalid_independent_one():
    """If pairing made no difference the comparison would not need it. It does."""
    p = pd.read_csv(TAB / "table65_vix_paired_comparison.csv")
    pair = p[p.quantity.str.contains("PAIRED bootstrap (used)", regex=False)].iloc[0]
    indep = p[p.quantity.str.contains("INDEPENDENT bootstrap", regex=False)].iloc[0]
    assert (pair.hi95 - pair.lo95) < (indep.hi95 - indep.lo95) / 2


def test_the_analytic_disagreement_is_disclosed_not_hidden():
    """Williams' t rejects where the bootstraps do not. That must be reported."""
    p = pd.read_csv(TAB / "table65_vix_paired_comparison.csv")
    will = p[p.quantity.str.contains("Williams", regex=False)].iloc[0]
    assert abs(float(will.value)) > 1.96
    assert "DISAGREES" in will.note
    import docx
    d = docx.Document(ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx")
    body = "\n".join(p_.text for p_ in d.paragraphs)
    assert "One check disagrees" in body


# ── M7 final: non-overlapping is not independent ─────────────────────────────────────────────

def test_forecast_origins_are_serially_dependent_so_blocking_is_required():
    """Non-overlapping windows remove the mechanical overlap, not the persistence. If this ever
    fell near zero the block bootstrap would be unnecessary -- and that should be a deliberate
    decision, not a silent one."""
    d = pd.read_csv(TAB / "table67_forward_origin_dependence.csv")
    assert float(d.lag1_autocorrelation.abs().max()) > 0.3, \
        "origins now look independent; revisit whether blocking is still needed"


def test_headline_forward_intervals_are_block_bootstrap_not_analytic():
    """The quoted interval must be the dependence-aware one. The analytic Fisher-z interval is
    roughly half as wide here, so quoting it would overstate precision."""
    t = pd.read_csv(TAB / "paper_table23_vix_forward_primary.csv")
    assert any("block bootstrap" in c for c in t.columns), \
        "Table 23 no longer quotes the paired block-bootstrap interval"
    bk = pd.read_csv(TAB / "table68_forward_block_bootstrap.csv").query("mean_block_origins == 3")
    an = pd.read_csv(TAB / "table59_vix_forward_primary.csv")
    for est in ("Close-to-close", "Parkinson"):
        w_blk = float(bk[bk.estimator == est].corr_hi95.iloc[0]
                      - bk[bk.estimator == est].corr_lo95.iloc[0])
        w_an = float(an[an.estimator == est].pearson_hi95.iloc[0]
                     - an[an.estimator == est].pearson_lo95.iloc[0])
        assert w_blk > w_an, f"{est}: the block interval is not wider than the analytic one"


def test_inconclusive_verdict_survives_every_block_length():
    db = pd.read_csv(TAB / "table69_forward_difference_block.csv")
    assert set(db.mean_block_origins) >= {1, 3, 6}
    assert db.contains_zero.all(), \
        "the estimator difference no longer contains zero at some block length"


def test_manuscript_does_not_claim_attenuation_for_close_to_close():
    """Its slope interval contains one under dependence-aware inference, so the claim that the
    slope is below one may not be made for that estimator."""
    import docx
    d = docx.Document(ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx")
    body = "\n".join(p.text for p in d.paragraphs)
    assert "Calibration slopes are well below one" not in body
    assert "contains one; we therefore do not assert attenuation" in body
    bk = pd.read_csv(TAB / "table68_forward_block_bootstrap.csv").query("mean_block_origins == 3")
    cc = bk[bk.estimator == "Close-to-close"].iloc[0]
    assert cc.slope_lo95 < 1.0 < cc.slope_hi95, \
        "close-to-close's slope interval no longer contains one; the wording must be revisited"


def test_analytic_fisher_helper_is_labelled_as_not_distribution_free():
    src = (ROOT / "scripts" / "32_vix_forward_validation.py").read_text()
    doc = src.split("def fisher_ci")[1].split("def ")[0]
    assert "NOT distribution-free" in doc and "NOT dependence-aware" in doc


# ── M7 close: the block length is selected, not chosen after seeing the intervals ─────────────

def test_primary_block_length_is_selected_from_the_data():
    """Reporting a grid and then quoting one member invites selection after the fact. The
    primary must be the automatic Politis-White choice, and it must be flagged in the outputs."""
    bl = pd.read_csv(TAB / "table70_block_length_selection.csv")
    assert (bl.politis_white_block_length > 1).any()
    prim = bl[bl.series.str.startswith("PRIMARY")]
    assert len(prim) == 1
    # the primary must equal the maximum over the component series
    comp = bl[~bl.series.str.startswith("PRIMARY")].politis_white_block_length.max()
    assert abs(float(prim.politis_white_block_length.iloc[0]) - comp) < 0.01
    db = pd.read_csv(TAB / "table69_forward_difference_block.csv")
    assert db.is_primary.sum() == 1, "exactly one block length must be flagged primary"


def test_origin_and_expected_block_counts_are_reported():
    """Readers need these to judge whether long blocks leave enough information. Note this is
    the EXPECTED NUMBER OF BLOCKS PER REPLICATE, not an effective sample size -- the two are
    different quantities and the manuscript must not conflate them."""
    db = pd.read_csv(TAB / "table69_forward_difference_block.csv")
    for col in ("n_origins", "expected_blocks_per_replicate"):
        assert col in db.columns
    prim = db[db.is_primary].iloc[0]
    assert int(prim.n_origins) > 100
    assert float(prim.expected_blocks_per_replicate) > 5


def test_max_aggregation_of_block_selectors_is_disclosed_as_our_own_rule():
    """Politis-White selects for ONE series. Taking the maximum across three is our choice,
    adopted post hoc, and must not be attributed to the procedure."""
    src = (ROOT / "scripts" / "32_vix_forward_validation.py").read_text()
    assert "CUSTOM AGGREGATION RULE OF OURS" in src
    assert "POST HOC" in src
    import docx
    d = docx.Document(ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx")
    body = "\n".join(p_.text for p_ in d.paragraphs)
    assert "our own" in body and "post hoc" in body
    assert "effectively independent blocks" not in body


def test_verdict_holds_at_the_selected_block_length():
    db = pd.read_csv(TAB / "table69_forward_difference_block.csv")
    assert bool(db[db.is_primary].contains_zero.iloc[0])
    assert db.contains_zero.all()


def test_manuscript_does_not_call_the_origins_independent_or_the_interval_distribution_free():
    import docx
    d = docx.Document(ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx")
    body = "\n".join(p.text for p in d.paragraphs)
    assert "independent forecast origins" not in body
    assert "distribution-free" not in body
    assert "after Fisher-z transformation" in body
    # the block length must be described as data-selected, however that is phrased
    assert "selected from the data" in body and "Politis-White" in body
    # and n/b must not be mislabelled as an effective sample size
    assert "expected-length blocks" in body
    assert "effectively independent blocks" not in body


# ── M13: the three literature tables must reconcile exactly ──────────────────────────────────

def _lit():
    return (pd.read_csv(TAB / "table71_nepal_search_log.csv"),
            pd.read_csv(TAB / "table72_nepal_literature_matrix.csv"),
            pd.read_csv(TAB / "table73_nepal_screening_ledger.csv"))


def test_literature_tables_reconcile():
    """located - duplicates - excluded - inaccessible = included = rows in the matrix."""
    _, matrix, ledger = _lit()
    located = len(ledger)
    dup = int((ledger.inclusion == "duplicate").sum())
    exc = int((ledger.inclusion == "excluded").sum())
    ina = int((ledger.inclusion == "inaccessible").sum())
    inc = int((ledger.inclusion == "included").sum())
    assert located - dup - exc - ina == inc
    assert inc == len(matrix)
    assert set(matrix.record_id) == set(ledger.record_id[ledger.inclusion == "included"])


def test_every_exclusion_carries_a_specific_reason():
    _, _, ledger = _lit()
    exc = ledger[ledger.inclusion == "excluded"]
    assert len(exc) > 0
    assert exc.exclusion_reason.fillna("").astype(str).str.strip().ne("").all()


def test_inaccessible_records_are_preserved_not_excluded():
    """An in-scope study we could not retrieve must stay visible in the ledger."""
    _, _, ledger = _lit()
    ina = ledger[ledger.inclusion == "inaccessible"]
    assert len(ina) > 0
    # empty cells read back as NaN, so normalise before checking
    reasons = ina.exclusion_reason.fillna("").astype(str).str.strip()
    assert (reasons == "").all(), \
        "an inaccessible record was given an exclusion reason, conflating the two categories"
    # and they must still be retrievable as a group, not folded into exclusions
    assert set(ina.retrieval_status) <= {"inaccessible", "not_retrieved"}


def test_duplicates_point_at_an_existing_canonical_record():
    _, _, ledger = _lit()
    dup = ledger[ledger.inclusion == "duplicate"]
    for _, r in dup.iterrows():
        assert str(r.duplicate_of).strip()
        assert r.duplicate_of in set(ledger.record_id)


def test_matrix_carries_the_ten_frozen_evidence_fields():
    _, matrix, _ = _lit()
    for f in ("citation", "doi_or_link", "study_period", "frequency", "level",
              "volatility_measure", "model_or_question", "thin_trading_treatment",
              "ohlc_range_estimators_used", "main_finding", "full_text_verified"):
        assert f in matrix.columns, f"the matrix lacks the frozen field {f!r}"


def test_full_text_verification_is_not_claimed_from_abstracts():
    """Rows verified only from a publisher record must not be marked full-text verified."""
    _, matrix, _ = _lit()
    for _, r in matrix.iterrows():
        if "abstract" in str(r.verification_level).lower():
            assert not bool(r.full_text_verified), \
                f"{r.record_id}: abstract-level verification marked as full text"


# ── M13: the manuscript's novelty language must stay inside the frozen matrix ─────────────────

BANNED_NOVELTY = ("primarily index-level", "fills a gap", "no prior study",
                  "first study", "the first to", "systematic review")


def _manuscript_body():
    import docx
    d = docx.Document(ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx")
    return "\n".join(p.text for p in d.paragraphs)


def test_no_banned_novelty_phrase_survives():
    body = _manuscript_body().lower()
    for phrase in BANNED_NOVELTY:
        assert phrase not in body, f"banned novelty phrase present: {phrase!r}"


def test_the_verified_security_level_study_is_cited():
    """Only R13 was full-text verified, and it is the study that retires the index-vs-security
    framing, so it must be named explicitly rather than lumped into a citation block."""
    body = _manuscript_body()
    assert "Neupane" in body
    matrix = pd.read_csv(TAB / "table72_nepal_literature_matrix.csv")
    verified = matrix[matrix.full_text_verified]
    assert len(verified) >= 1
    for surname in ("Neupane",):
        assert surname in body


def test_abstract_only_studies_are_not_used_to_establish_absence():
    """Dangal & Gajurel and Karki are verified to abstract level only. They may be cited for
    what their abstracts say and nothing more."""
    body = _manuscript_body()
    for cite in ("Dangal", "Karki"):
        assert cite in body
    # no absence/novelty claim may sit in the same sentence as either citation
    for sentence in re.split(r"(?<=[.!?])\s+", body):
        if "Dangal" in sentence or "Karki" in sentence:
            low = sentence.lower()
            for bad in ("no prior", "first", "never been", "has not been", "unlike any"):
                assert bad not in low, \
                    f"an absence/novelty claim sits with an abstract-only citation: {sentence!r}"


def test_contribution_is_stated_as_measurement_not_as_priority():
    body = _manuscript_body()
    assert "no claim is made to introduce security-level" in body
    assert "makes no claim to introduce security-level" in body
    assert "addresses a measurement question not identified in the targeted search" in body


def test_abstract_stays_within_the_250_word_cap():
    import docx
    d = docx.Document(ROOT / "paper" / "NEPSE_Volatility_Manuscript_Revised_2026-09.docx")
    ps = d.paragraphs
    i = next(k for k, p in enumerate(ps) if p.text.strip() == "Abstract")
    words = sum(len(p.text.split()) for p in ps[i + 1:i + 5])
    assert words <= 255, f"abstract is {words} words including its four labels"


def test_search_quality_control_stays_out_of_the_manuscript():
    body = _manuscript_body().lower()
    for term in ("snippet", "search summary", "26 securities", "misattribut"):
        assert term not in body, f"search-quality control leaked into the manuscript: {term!r}"
